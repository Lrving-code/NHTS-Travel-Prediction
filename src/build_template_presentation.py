"""Build a dense, template-inspired academic presentation deck."""

from __future__ import annotations

from io import BytesIO
import logging
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = PROJECT_ROOT / "temp" / "ppt_template_reference.pptx"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "final_project" / "NHTS_Travel_Behavior_Template_Presentation.pptx"
FINAL_DIR = PROJECT_ROOT / "outputs" / "final_project"
MODE_DIR = PROJECT_ROOT / "outputs" / "mode_composition_extension"
PURPOSE_DIR = PROJECT_ROOT / "outputs" / "purpose_composition_extension"
STATS_DIR = PROJECT_ROOT / "outputs" / "statistical_validation"
LLM_FEATURE_DIR = PROJECT_ROOT / "outputs" / "llm_event_features"
PARETO_DIR = PROJECT_ROOT / "outputs" / "multi_objective_pareto"
FIGURE_DIR = FINAL_DIR / "figures" / "presentation_figures"
ROBUSTNESS_DIR = PROJECT_ROOT / "outputs" / "robustness_checks"
COHORT_PRIOR_DIR = PROJECT_ROOT / "outputs" / "cohort_prior_value_analysis"
ZERO_SHOT_RULE_TREE_PATH = PROJECT_ROOT / "outputs" / "zero_shot_llm_rule_tree_baseline" / "zero_shot_llm_rule_tree_metrics.csv"
SMALL_DATA_CALIBRATION_PATH = (
    PROJECT_ROOT / "outputs" / "llm_rule_small_data_calibration" / "method_spectrum_metrics.csv"
)
EMU_PER_INCH = 914400
TOTAL_SLIDES = 24

LOGGER = logging.getLogger(__name__)

COLORS = {
    "navy": RGBColor(17, 42, 66),
    "navy2": RGBColor(24, 54, 80),
    "teal": RGBColor(0, 121, 107),
    "green": RGBColor(25, 135, 84),
    "blue": RGBColor(37, 99, 235),
    "orange": RGBColor(234, 88, 12),
    "red": RGBColor(190, 18, 60),
    "purple": RGBColor(91, 33, 182),
    "gray": RGBColor(100, 116, 139),
    "line": RGBColor(203, 213, 225),
    "pale": RGBColor(241, 245, 249),
    "light": RGBColor(248, 250, 252),
    "ink": RGBColor(15, 23, 42),
    "white": RGBColor(255, 255, 255),
}

METHOD_LABELS = {
    "historical_mean_only": "Historical mean",
    "historical_xgboost": "Ordinary XGBoost",
    "historical_mean_trend_shift": "Trend / indicator",
    "llm_only_trip_suppression_a1p25": "LLM-only pressure",
    "zero_shot_llm_rule_tree": "Zero-shot rule tree",
    "llm_rule_small_hist_calibrated_n500": "Rule + 500 history",
    "global_trip_suppression_a1p25": "Global event rule",
    "random_trip_suppression_a1p25": "Random pressure",
    "gated_trip_suppression_a1_d0p15": "Hybrid gated",
}


def configure_logging() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")


def clear_template_slides(prs: Presentation) -> None:
    slide_id_list = prs.slides._sldIdLst  # noqa: SLF001 - python-pptx has no public clear API.
    for slide_id in list(slide_id_list):
        prs.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def iter_nested_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_nested_shapes(shape.shapes)


def extract_template_logo() -> bytes | None:
    if not TEMPLATE_PATH.exists():
        return None
    prs = Presentation(TEMPLATE_PATH)
    for slide in prs.slides:
        for shape in iter_nested_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            width = shape.width / EMU_PER_INCH
            height = shape.height / EMU_PER_INCH
            size = len(shape.image.blob)
            if 0.7 <= width <= 1.4 and 0.35 <= height <= 0.95 and size < 80_000:
                return shape.image.blob
    return None


def blank_slide(prs: Presentation) -> Slide:
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_background(slide: Slide, color: RGBColor = COLORS["white"]) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def rect(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    line: RGBColor | None = None,
    radius: bool = False,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def text_box(
    slide: Slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 10,
    color: RGBColor = COLORS["ink"],
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    for run in paragraph.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_logo(slide: Slide, logo: bytes | None, x: float, y: float, w: float, dark: bool = False) -> None:
    if logo is None:
        text_box(slide, "TSINGHUA", x - 0.05, y + 0.05, w + 0.15, 0.22, 7, COLORS["gray"], True, PP_ALIGN.CENTER)
        return
    if dark:
        rect(slide, x - 0.04, y - 0.03, w + 0.08, 0.62, COLORS["white"], None, True)
    slide.shapes.add_picture(BytesIO(logo), Inches(x), Inches(y), width=Inches(w))


def add_picture(slide: Slide, path: Path, x: float, y: float, w: float) -> None:
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        rect(slide, x, y, w, 2.6, COLORS["pale"], COLORS["line"], True)
        text_box(slide, f"Missing figure:\n{path.name}", x + 0.1, y + 1.0, w - 0.2, 0.5, 12, COLORS["red"], True, PP_ALIGN.CENTER)


def story_box(slide: Slide, zh: str, en: str, x: float, y: float, w: float, h: float, color: RGBColor) -> None:
    rect(slide, x, y, w, h, COLORS["pale"], COLORS["line"], True)
    rect(slide, x, y, 0.055, h, color)
    zh_height = min(0.46, max(0.32, h * 0.36))
    text_box(slide, zh, x + 0.18, y + 0.1, w - 0.32, zh_height, 12, COLORS["ink"], True, valign=MSO_ANCHOR.TOP)
    text_box(
        slide,
        en,
        x + 0.18,
        y + 0.17 + zh_height,
        w - 0.32,
        max(0.34, h - zh_height - 0.25),
        12,
        COLORS["gray"],
        valign=MSO_ANCHOR.TOP,
    )


def add_frame(
    slide: Slide,
    section: str,
    title: str,
    subtitle: str,
    page: int,
    logo: bytes | None,
    dark: bool = False,
) -> None:
    title_color = COLORS["white"] if dark else COLORS["ink"]
    sub_color = RGBColor(203, 213, 225) if dark else COLORS["gray"]
    line_color = RGBColor(148, 163, 184) if dark else COLORS["line"]
    text_box(slide, "Mobility Inequality and Sustainable Development", 0.42, 0.12, 4.9, 0.18, 7.5, sub_color, True)
    text_box(slide, section, 0.42, 0.38, 0.55, 0.32, 14, COLORS["teal"], True)
    text_box(slide, title, 1.25, 0.27, 9.2, 0.36, 17.5, title_color, True)
    if subtitle:
        text_box(slide, subtitle, 1.27, 0.64, 9.7, 0.25, 8.5, sub_color)
    add_logo(slide, logo, 12.05, 0.22, 0.78, dark)
    rect(slide, 0.42, 0.96, 12.25, 0.012, line_color)
    rect(slide, 0.42, 7.03, 12.25, 0.012, line_color)
    text_box(slide, "LLM-Guided Temporal Adaptation for Household Mobility", 0.42, 7.08, 6.25, 0.18, 7.5, sub_color)
    text_box(slide, f"{page:02d} / {TOTAL_SLIDES}", 11.65, 7.08, 1.0, 0.18, 7.5, sub_color, True, PP_ALIGN.RIGHT)


def add_backup_frame(slide: Slide, title: str, subtitle: str, label: str, logo: bytes | None) -> None:
    text_box(slide, "Mobility Inequality and Sustainable Development", 0.42, 0.12, 4.9, 0.18, 7.5, COLORS["gray"], True)
    text_box(slide, "BACKUP", 0.42, 0.38, 0.85, 0.32, 13, COLORS["orange"], True)
    text_box(slide, title, 1.35, 0.27, 9.2, 0.36, 17.0, COLORS["ink"], True)
    if subtitle:
        text_box(slide, subtitle, 1.37, 0.64, 9.7, 0.25, 8.5, COLORS["gray"])
    add_logo(slide, logo, 12.05, 0.22, 0.78)
    rect(slide, 0.42, 0.96, 12.25, 0.012, COLORS["line"])
    rect(slide, 0.42, 7.03, 12.25, 0.012, COLORS["line"])
    text_box(slide, "Appendix: reviewer and instructor Q&A", 0.42, 7.08, 5.7, 0.18, 7.5, COLORS["gray"])
    text_box(slide, label, 11.1, 7.08, 1.55, 0.18, 7.5, COLORS["gray"], True, PP_ALIGN.RIGHT)


def metric_card(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    value: str,
    note: str,
    color: RGBColor,
    fill: RGBColor = COLORS["light"],
) -> None:
    rect(slide, x, y, w, h, fill, COLORS["line"], True)
    rect(slide, x, y, 0.055, h, color)
    text_box(slide, label, x + 0.14, y + 0.08, w - 0.24, 0.18, 7.8, COLORS["gray"], True)
    text_box(slide, value, x + 0.14, y + 0.27, w - 0.24, 0.32, 16, color, True)
    text_box(slide, note, x + 0.14, y + 0.61, w - 0.24, 0.2, 7.5, COLORS["gray"])


def bullet_list(slide: Slide, items: list[str], x: float, y: float, w: float, h: float, size: float = 10) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {item}"
        paragraph.level = 0
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = COLORS["ink"]
        paragraph.space_after = Pt(2)


def node(
    slide: Slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    note: str,
    color: RGBColor,
    fill: RGBColor = COLORS["white"],
) -> None:
    rect(slide, x, y, w, h, fill, color, True)
    text_box(slide, title, x + 0.08, y + 0.07, w - 0.16, 0.2, 9.2, color, True, PP_ALIGN.CENTER)
    text_box(slide, note, x + 0.1, y + 0.31, w - 0.2, h - 0.37, 7.8, COLORS["gray"], False, PP_ALIGN.CENTER)


def arrow(slide: Slide, x: float, y: float, w: float, h: float, color: RGBColor = COLORS["gray"]) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def small_table(
    slide: Slide,
    headers: list[str],
    rows: list[list[str]],
    x: float,
    y: float,
    widths: list[float],
    row_h: float = 0.34,
    font_size: float = 7.7,
    header_color: RGBColor = COLORS["navy"],
) -> None:
    actual_size = max(font_size, 12)
    actual_row_h = max(row_h, 0.42)
    for col, header in enumerate(headers):
        xx = x + sum(widths[:col])
        rect(slide, xx, y, widths[col], actual_row_h, header_color)
        text_box(
            slide,
            header,
            xx + 0.04,
            y + 0.08,
            widths[col] - 0.08,
            actual_row_h - 0.14,
            actual_size,
            COLORS["white"],
            True,
            PP_ALIGN.CENTER,
        )
    for row_idx, row in enumerate(rows):
        yy = y + actual_row_h * (row_idx + 1)
        fill = COLORS["light"] if row_idx % 2 == 0 else COLORS["pale"]
        for col, value in enumerate(row):
            xx = x + sum(widths[:col])
            rect(slide, xx, yy, widths[col], actual_row_h, fill, COLORS["line"])
            text_box(
                slide,
                value,
                xx + 0.06,
                yy + 0.07,
                widths[col] - 0.12,
                actual_row_h - 0.14,
                actual_size,
                COLORS["ink"],
                col == 0,
                valign=MSO_ANCHOR.TOP,
            )


def method_comparison_table(
    slide: Slide,
    rows: list[tuple[str, bool, bool, bool, str, str]],
    x: float,
    y: float,
    widths: list[float],
    row_h: float = 0.36,
) -> None:
    actual_row_h = max(row_h, 0.38)
    table_font = 10.5 if actual_row_h < 0.45 else 12
    headers = ["Method", "传统基线", "LLM先验", "目标年标签", "对比作用", "wMAE"]
    for col, header in enumerate(headers):
        xx = x + sum(widths[:col])
        rect(slide, xx, y, widths[col], actual_row_h, COLORS["navy"])
        text_box(slide, header, xx + 0.04, y + 0.07, widths[col] - 0.08, actual_row_h - 0.12, table_font, COLORS["white"], True, PP_ALIGN.CENTER)
    for row_idx, (method, historical, llm_prior, target_label, role, value) in enumerate(rows):
        yy = y + actual_row_h * (row_idx + 1)
        fill = COLORS["light"] if row_idx % 2 == 0 else COLORS["pale"]
        values = [method, historical, llm_prior, target_label, role, value]
        for col, cell in enumerate(values):
            xx = x + sum(widths[:col])
            rect(slide, xx, yy, widths[col], actual_row_h, fill, COLORS["line"])
            if col in {1, 2, 3}:
                symbol = "●" if cell else "×"
                color = COLORS["green"] if cell else COLORS["red"]
                text_box(slide, symbol, xx + 0.05, yy + 0.06, widths[col] - 0.1, actual_row_h - 0.12, 12, color, True, PP_ALIGN.CENTER)
            else:
                text = str(cell)
                align = PP_ALIGN.RIGHT if col == 5 else PP_ALIGN.LEFT
                bold = col == 0 or method == "Hybrid gated"
                text_box(slide, text, xx + 0.06, yy + 0.06, widths[col] - 0.12, actual_row_h - 0.12, table_font, COLORS["ink"], bold, align)


def draw_bar_chart(
    slide: Slide,
    values: list[tuple[str, float, RGBColor]],
    x: float,
    y: float,
    w: float,
    h: float,
    max_value: float,
    unit: str = "",
) -> None:
    row_h = h / len(values)
    for index, (label, value, color) in enumerate(values):
        yy = y + index * row_h
        text_box(slide, label, x, yy - 0.015, 2.25, 0.24, 7.8, COLORS["ink"])
        rect(slide, x + 2.35, yy + 0.04, w - 3.0, 0.17, RGBColor(226, 232, 240))
        width = (w - 3.0) * min(value / max_value, 1.0)
        rect(slide, x + 2.35, yy + 0.04, width, 0.17, color)
        text_box(slide, f"{value:.2f}{unit}", x + w - 0.55, yy - 0.02, 0.55, 0.24, 8, color, True, PP_ALIGN.RIGHT)


def draw_bias_strip(slide: Slide, values: list[tuple[str, float, RGBColor]], x: float, y: float, w: float, h: float) -> None:
    center = x + w * 0.52
    rect(slide, center, y, 0.015, h, COLORS["gray"])
    text_box(slide, "0", center - 0.12, y + h + 0.02, 0.24, 0.15, 7, COLORS["gray"], True, PP_ALIGN.CENTER)
    scale = (w * 0.45) / 4.8
    for idx, (label, value, color) in enumerate(values):
        yy = y + 0.18 + idx * 0.48
        text_box(slide, label, x, yy - 0.03, 1.55, 0.22, 7.5, COLORS["ink"])
        bar_w = abs(value) * scale
        if value >= 0:
            bx = center
        else:
            bx = center - bar_w
        rect(slide, bx, yy, bar_w, 0.15, color)
        text_box(slide, f"{value:+.2f}", bx + (bar_w if value >= 0 else -0.5), yy - 0.04, 0.5, 0.24, 7.5, color, True)


def read_csv_or_empty(path: Path) -> pd.DataFrame:
    if path.exists():
        return pd.read_csv(path)
    return pd.DataFrame()


def append_optional_trip_rows(trip: pd.DataFrame) -> pd.DataFrame:
    optional_frames: list[pd.DataFrame] = []
    if ZERO_SHOT_RULE_TREE_PATH.exists():
        zero = pd.read_csv(ZERO_SHOT_RULE_TREE_PATH)
        zero = zero[zero["method"].isin(["zero_shot_llm_rule_tree"])]
        optional_frames.append(zero)
    if SMALL_DATA_CALIBRATION_PATH.exists():
        small = pd.read_csv(SMALL_DATA_CALIBRATION_PATH)
        small = small[small["method"].isin(["llm_rule_small_hist_calibrated_n500"])]
        optional_frames.append(small)
    if not optional_frames:
        return trip

    combined = pd.concat([trip, *optional_frames], ignore_index=True, sort=False)
    combined = combined.drop_duplicates(subset=["method"], keep="last")
    return combined


def load_results() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    trip = append_optional_trip_rows(pd.read_csv(FINAL_DIR / "final_metrics_summary.csv"))
    acc = pd.read_csv(FINAL_DIR / "household_accuracy_summary.csv")
    mode = pd.read_csv(MODE_DIR / "mode_composition_metrics.csv")
    mode_trips = read_csv_or_empty(MODE_DIR / "mode_specific_trip_count_metrics.csv")
    purpose = read_csv_or_empty(PURPOSE_DIR / "purpose_composition_metrics.csv")
    ci = read_csv_or_empty(STATS_DIR / "trip_metric_confidence_intervals.csv")
    paired_ci = read_csv_or_empty(STATS_DIR / "paired_improvement_confidence_intervals.csv")
    return trip, acc, mode, mode_trips, purpose, ci, paired_ci


def metric(df: pd.DataFrame, method: str, column: str) -> float:
    return float(df.loc[df["method"] == method, column].iloc[0])


def pct_delta(before: float, after: float) -> float:
    return (before - after) / before


def ci_values(ci: pd.DataFrame, method: str, metric_name: str) -> tuple[float, float, float]:
    if ci.empty:
        return float("nan"), float("nan"), float("nan")
    row = ci[(ci["method"] == method) & (ci["metric"] == metric_name)].iloc[0]
    return float(row["point"]), float(row["ci95_low"]), float(row["ci95_high"])


def paired_ci_value(paired: pd.DataFrame, metric_name: str) -> tuple[float, float, float]:
    if paired.empty:
        return float("nan"), float("nan"), float("nan")
    row = paired[paired["metric"] == metric_name].iloc[0]
    return float(row["mean"]), float(row["ci95_low"]), float(row["ci95_high"])


def cohort_value(quantity: str) -> float:
    summary_path = COHORT_PRIOR_DIR / "cohort_prior_value_summary.csv"
    if not summary_path.exists():
        return float("nan")
    summary = pd.read_csv(summary_path)
    rows = summary.loc[summary["quantity"] == quantity, "value"]
    if rows.empty:
        return float("nan")
    return float(rows.iloc[0])


def add_title_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide, COLORS["navy"])
    add_logo(slide, logo, 11.82, 0.34, 0.92, True)
    text_box(slide, "面向时序迁移的家庭出行预测\n模拟选择器与大模型修正框架", 0.68, 0.78, 9.9, 0.95, 24, COLORS["white"], True)
    text_box(slide, "LLM-Guided Temporal Adaptation for Household Mobility Prediction", 0.72, 1.86, 10.2, 0.32, 11.5, RGBColor(203, 213, 225), True)
    text_box(slide, "prediction selector + LLM-corrected XGBoost；不使用目标年出行标签校准", 0.72, 2.23, 9.7, 0.28, 10, RGBColor(203, 213, 225))
    rect(slide, 0.72, 3.0, 11.9, 0.01, RGBColor(148, 163, 184))
    nodes = [
        ("NHTS 历史数据", "routine mobility"),
        ("目标年情境变化", "temporal shift"),
        ("模拟预测选择器", "prediction selector"),
        ("LLM 修正 XGBoost", "corrected predictor"),
    ]
    for idx, (title, note) in enumerate(nodes):
        node(slide, 0.8 + idx * 3.02, 3.35, 2.25, 0.7, title, note, COLORS["teal"], RGBColor(30, 64, 91))
        if idx < len(nodes) - 1:
            arrow(slide, 3.1 + idx * 3.02, 3.57, 0.48, 0.18, RGBColor(148, 163, 184))
    metric_card(slide, 0.82, 5.05, 2.55, 0.9, "Trip-count MAE", f"-{values['mae_reduction']:.1%}", "vs ordinary XGBoost", COLORS["blue"], COLORS["white"])
    metric_card(slide, 3.68, 5.05, 2.55, 0.9, "Weighted bias", f"{values['gated_bias']:+.3f}", "nearly unbiased", COLORS["green"], COLORS["white"])
    metric_card(slide, 6.54, 5.05, 2.55, 0.9, "Within 2 trips", f"{values['within2']:.1%}", "household tolerance", COLORS["orange"], COLORS["white"])
    metric_card(slide, 9.4, 5.05, 2.55, 0.9, "Behavior scope", "3+ outputs", "trips / modes / purposes", COLORS["purple"], COLORS["white"])
    text_box(slide, "Mobility Inequality and Sustainable Development", 0.72, 6.83, 4.9, 0.2, 8, RGBColor(203, 213, 225), True)
    text_box(slide, f"01 / {TOTAL_SLIDES}", 11.72, 6.83, 0.85, 0.2, 8, RGBColor(203, 213, 225), True, PP_ALIGN.RIGHT)


def add_content_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "00", "目录 / Content", "从研究问题到技术路线，再到结果解释", 2, logo)
    rect(slide, 0.55, 1.45, 3.1, 4.7, COLORS["navy"])
    text_box(slide, "CONTENT", 1.02, 2.45, 2.1, 0.42, 20, COLORS["white"], True, PP_ALIGN.CENTER)
    text_box(slide, "10 分钟主线：\n为什么目标年需要时序适应，LLM 如何提供可泛化的情境先验。", 0.92, 3.1, 2.38, 0.9, 10, RGBColor(203, 213, 225), False, PP_ALIGN.CENTER)
    items = [
        ("01", "Background & gap", "时序迁移、移动不平等与相关研究缺口"),
        ("02", "Problem & data", "家庭层面多输出 travel behavior system"),
        ("03", "Method", "模拟预测选择器 + LLM-corrected XGBoost 的双分支框架"),
        ("04", "Evaluation", "性能比较、Pareto 选择、置信区间、稳健性检验"),
        ("05", "Insights", "出行强度、方式、目的与未来工作"),
    ]
    for idx, (num, title, note) in enumerate(items):
        y = 1.32 + idx * 0.94
        rect(slide, 4.12, y, 0.5, 0.5, COLORS["teal"], None, True)
        text_box(slide, num, 4.17, y + 0.08, 0.4, 0.22, 11, COLORS["white"], True, PP_ALIGN.CENTER)
        text_box(slide, title, 4.82, y + 0.02, 3.1, 0.22, 12, COLORS["ink"], True)
        text_box(slide, note, 4.82, y + 0.28, 6.8, 0.2, 8.5, COLORS["gray"])
        rect(slide, 4.55, y + 0.58, 7.4, 0.008, COLORS["line"])


def add_problem_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "01", "研究问题 / Research Gap", "目标年不是普通跨年预测 | Target-year shift is not ordinary transfer", 3, logo)
    headers = ["Traditional framing", "Event-driven framing"]
    rows = [
        ["2017 -> 2022 as smooth temporal transfer", "2022 as target-year behavioral regime shift"],
        ["Covariates explain most household differences", "Household covariates miss remote work / transit avoidance"],
        ["Model failure means weak tabular learner", "Model failure reveals unobserved event mechanisms"],
        ["Tune on target labels if available", "Use no 2022 labels for training/calibration"],
    ]
    small_table(slide, headers, rows, 0.68, 1.28, [5.55, 5.85], 0.38, 8.3)
    metric_card(slide, 0.72, 4.75, 2.6, 0.86, "Ordinary XGBoost", f"{values['baseline']:.2f}", "weighted MAE", COLORS["red"])
    metric_card(slide, 3.58, 4.75, 2.6, 0.86, "Systematic bias", f"{values['baseline_bias']:+.2f}", "over-predicts trips", COLORS["red"])
    metric_card(slide, 6.44, 4.75, 2.6, 0.86, "Pre-COVID bias", "0.90", "mean abs. wBias", COLORS["teal"])
    metric_card(slide, 9.3, 4.75, 2.6, 0.86, "Research target", "adapt", "not just fit", COLORS["blue"])
    story_box(
        slide,
        "核心问题：2022 的 weighted bias 远高于常规跨年验证；不使用 2022 标签调参时，能否用情境先验修正这种系统性高估。",
        "Core question: can context priors correct target-year overprediction without using 2022 target labels for calibration?",
        0.82,
        5.72,
        10.95,
        1.18,
        COLORS["teal"],
    )


def add_related_work_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "02", "相关研究与缺口 / Related Work and Gap", "2025-2026: LLM mobility, causal events, and foundation models", 4, logo)
    headers = ["最新方向", "2025-2026 代表线索", "本项目的位置"]
    rows = [
        ["Event-driven mobility", "ELLMob ICLR'26; CausalMob KDD'25", "把外部情境变化转成 label-free event priors"],
        ["Zero-shot LLM mobility", "AgentMove NAACL'25", "直接 LLM 可做 baseline，但数值校准弱"],
        ["Evidence-grounded agents", "AgentMob 2026", "routine fast path + low-confidence LLM reasoning"],
        ["Efficient LLM pipeline", "ELP-Mob GIS'25", "batch prompting 和调用压缩是必要工程约束"],
        ["Mobility foundation models", "UniMob KDD'25; STFM survey'25", "不训练大 FM；用轻量 adapter 处理 survey shift"],
        ["Planning evaluation", "causal guardrails + multi-objective metrics", "同时报告 accuracy、bias、equity、mode/purpose"],
    ]
    small_table(slide, headers, rows, 0.65, 1.22, [2.55, 4.2, 4.7], 0.50, 12, COLORS["navy2"])
    rect(slide, 0.9, 5.25, 10.8, 0.85, COLORS["pale"], COLORS["line"], True)
    text_box(slide, "研究缺口", 1.12, 5.52, 1.1, 0.22, 10, COLORS["orange"], True)
    text_box(
        slide,
        "缺口：已有工作多关注 trajectory / flow / sensor forecasting；NHTS household survey 的 target-year label-free temporal adaptation 仍需要一个可解释、可审计框架。",
        2.25,
        5.37,
        8.95,
        0.42,
        12,
        COLORS["ink"],
        True,
    )


def add_data_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "03", "数据与目标 / Data and Target System", "一个 household behavior system，多种 linked outputs", 5, logo)
    years = [("2001", "history"), ("2009", "history"), ("2017", "history + mode base"), ("2022", "target-year eval")]
    for idx, (year, label) in enumerate(years):
        x = 0.75 + idx * 2.65
        rect(slide, x, 1.52, 1.0, 0.55, COLORS["teal"] if year != "2022" else COLORS["orange"], None, True)
        text_box(slide, year, x + 0.08, 1.62, 0.84, 0.22, 12, COLORS["white"], True, PP_ALIGN.CENTER)
        text_box(slide, label, x - 0.25, 2.13, 1.5, 0.22, 8, COLORS["gray"], False, PP_ALIGN.CENTER)
        if idx < len(years) - 1:
            arrow(slide, x + 1.15, 1.72, 0.72, 0.16, COLORS["gray"])
    headers = ["Output", "Construction", "Planning meaning", "Main metric"]
    rows = [
        ["Trip generation", "CNTTDHH at household level", "How much travel demand", "weighted MAE / bias"],
        ["Mode composition", "TRPTRANS -> household share vector", "How demand is distributed", "TV / share MAE"],
        ["Purpose composition", "TRIPPURP -> household purpose vector", "Why households travel", "TV / purpose MAE"],
        ["Mode-specific trips", "predicted trips × mode share", "Trips by mode for planning", "derived output"],
    ]
    small_table(slide, headers, rows, 0.72, 3.02, [2.35, 3.5, 3.05, 2.35], 0.42, 8.3, COLORS["navy2"])
    story_box(
        slide,
        "目标不是只预测 CNTTDHH，而是构造 household travel behavior：出行多少次、用什么方式、为什么出行、每种方式多少次。",
        "The target is a household behavior system: trip generation, mode composition, purpose composition, and derived mode-specific trip counts.",
        0.82,
        5.55,
        10.95,
        1.12,
        COLORS["blue"],
    )


def add_route_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "05", "技术路线 / Technical Route", "selector branch 处理冷启动，correction branch 处理目标年迁移", 7, logo)
    col_x = [1.0, 3.65, 6.3, 8.95]
    box_w = 2.05
    gap_arrow_w = 0.38

    text_box(slide, "Historical routine branch", 0.72, 1.18, 2.35, 0.22, 9.5, COLORS["blue"], True)
    for idx, (title, note) in enumerate(
        [
            ("NHTS history", "2001 / 2009 / 2017"),
            ("Feature harmonization", "same household schema"),
            ("XGBoost CUDA", "learn f_hist(X)"),
        ]
    ):
        node(slide, col_x[idx], 1.48, box_w, 0.68, title, note, COLORS["blue"])
        if idx < 2:
            arrow(slide, col_x[idx] + box_w + 0.08, 1.74, gap_arrow_w, 0.14, COLORS["gray"])

    text_box(slide, "LLM context-prior / selector branch", 0.72, 2.58, 3.25, 0.22, 9.5, COLORS["orange"], True)
    for idx, (title, note) in enumerate(
        [
            ("Target-year profile", "aggregated households"),
            ("GPT-5.5 prompt", "context + rule signals"),
            ("Validated priors", "s_event + selector"),
        ]
    ):
        node(slide, col_x[idx], 2.88, box_w, 0.68, title, note, COLORS["orange"])
        if idx < 2:
            arrow(slide, col_x[idx] + box_w + 0.08, 3.14, gap_arrow_w, 0.14, COLORS["gray"])

    rect(slide, 0.72, 4.08, 7.65, 1.18, COLORS["pale"], COLORS["line"], True)
    text_box(slide, "Two adaptation branches", 0.95, 4.23, 2.25, 0.22, 10, COLORS["green"], True)
    text_box(slide, "ŷ_2022 = f_hist(X) × clip(1 − α × s_event, min_factor, 1)", 2.75, 4.16, 5.25, 0.48, 12, COLORS["ink"], True)
    text_box(slide, "selector branch: LLM rules -> confidence gate -> optional fallback", 2.75, 4.61, 5.25, 0.28, 9.2, COLORS["gray"])
    text_box(slide, "target-year labels enter only here: final evaluation", 2.75, 4.93, 5.25, 0.28, 9.2, COLORS["red"], True)
    arrow(slide, 6.7, 2.15, 0.16, 0.62, COLORS["gray"])
    arrow(slide, 6.7, 3.56, 0.16, 0.44, COLORS["gray"])

    text_box(slide, "Output behavior system", 9.0, 4.1, 2.5, 0.22, 10, COLORS["green"], True)
    node(slide, 9.0, 4.45, 2.05, 0.62, "Trip generation", "household CNTTDHH", COLORS["green"])
    node(slide, 11.0, 4.45, 1.6, 0.62, "Mode shares", "transit signal", COLORS["green"])
    arrow(slide, 8.42, 4.58, 0.42, 0.16, COLORS["green"])

    rect(slide, 9.0, 1.35, 3.05, 2.6, COLORS["light"], COLORS["line"], True)
    text_box(slide, "Leakage guardrails", 9.25, 1.55, 2.55, 0.22, 10.5, COLORS["red"], True, PP_ALIGN.CENTER)
    small_table(
        slide,
        ["禁止进入 LLM", "原因"],
        [
            ["CNTTDHH", "目标标签"],
            ["WTHHFIN", "样本权重"],
            ["HOUSEID", "个体标识"],
        ],
        9.2,
        1.92,
        [1.35, 1.35],
        0.34,
        12,
        COLORS["red"],
    )


def add_llm_prior_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "06", "LLM 情境先验 / LLM Context Prior", "从 cohort 描述生成结构化目标年响应变量", 8, logo)
    headers = ["LLM output", "变量含义", "进入模型的位置"]
    rows = [
        ["trip_suppression", "总体出行折减压力", "修正 total trips"],
        ["remote_work", "通勤被远程办公替代的可能性", "解释出行下降来源"],
        ["transit_avoidance", "公共交通规避倾向", "修正 transit share"],
        ["delivery_substitution", "购物/服务被线上替代", "解释非通勤出行下降"],
        ["recovery_sensitivity", "恢复到常态出行的能力", "刻画 rebound 差异"],
        ["confidence", "LLM 对判断的自评置信度", "质量检查信号"],
    ]
    small_table(slide, headers, rows, 0.68, 1.18, [2.35, 4.45, 4.3], 0.36, 8.1, COLORS["navy2"])
    small_table(
        slide,
        ["LLM 输入", "明确排除"],
        [
            ["cohort-level household profile", "CNTTDHH / trip count label"],
            ["urban context, vehicles, workers", "sample weights / household IDs"],
            ["generic 2022 pandemic context", "aggregate 2022 outcomes"],
        ],
        0.82,
        4.85,
        [5.15, 5.15],
        0.34,
        8.2,
        COLORS["teal"],
    )

def add_comparison_slide(prs: Presentation, logo: bytes | None, trip: pd.DataFrame) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "07", "对比体系 / Evaluation Design", "传统模型、selector、纯 LLM 与 hybrid correction 逐层比较", 9, logo)
    rows = [
        ("Historical mean", True, False, False, "只保留历史平均水平", f"{metric(trip, 'historical_mean_only', 'weighted_mae'):.2f}"),
        ("Ordinary XGBoost", True, False, False, "传统方法：家庭属性 -> 出行次数", f"{metric(trip, 'historical_xgboost', 'weighted_mae'):.2f}"),
        ("Trend / indicator", True, False, False, "统一趋势修正，不区分情境机制", f"{metric(trip, 'historical_mean_trend_shift', 'weighted_mae'):.2f}"),
        ("LLM-only pressure", False, True, False, "只有事件方向，缺少 household baseline", f"{metric(trip, 'llm_only_trip_suppression_a1p25', 'weighted_mae'):.2f}"),
        ("Zero-shot rule tree", False, True, False, "LLM 直接构建规则树，无数值校准", f"{metric(trip, 'zero_shot_llm_rule_tree', 'weighted_mae'):.2f}"),
        ("Rule + 500 history", True, True, False, "LLM 规则 + 少量历史标签校准", f"{metric(trip, 'llm_rule_small_hist_calibrated_n500', 'weighted_mae'):.2f}"),
        ("Global event rule", True, False, False, "所有家庭使用同一目标年折减", f"{metric(trip, 'global_trip_suppression_a1p25', 'weighted_mae'):.2f}"),
        ("Random pressure", True, False, False, "随机 pressure 对照", f"{metric(trip, 'random_trip_suppression_a1p25', 'weighted_mae'):.2f}"),
        ("Hybrid gated", True, True, False, "历史基线 × LLM 事件折减", f"{metric(trip, 'gated_trip_suppression_a1_d0p15', 'weighted_mae'):.2f}"),
    ]
    method_comparison_table(slide, rows, 0.52, 1.12, [2.15, 1.0, 0.9, 0.95, 4.8, 0.75], 0.39)
    text_box(slide, "● 使用 / used    × 不使用 / not used", 0.65, 5.18, 3.7, 0.24, 12, COLORS["gray"], True)
    story_box(
        slide,
        "对比逻辑：LLM 直接构树能给出事件方向，但缺少历史行为的数值锚点；小样本规则校准仍偏高估。主方法保留 household baseline，再用 LLM 事件先验做无标签修正。",
        "The comparison isolates household grounding, zero-shot rule priors, small historical calibration, and the final gated event adapter.",
        0.82,
        5.45,
        10.95,
        1.14,
        COLORS["teal"],
    )


def add_results_slide(prs: Presentation, logo: bytes | None, trip: pd.DataFrame, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "08", "主结果 / Main Trip-Count Results", "Hybrid adapters reduce error while controlling bias", 10, logo)
    chart_values = [
        (METHOD_LABELS["historical_mean_only"], metric(trip, "historical_mean_only", "weighted_mae"), COLORS["gray"]),
        (METHOD_LABELS["historical_xgboost"], values["baseline"], COLORS["red"]),
        (METHOD_LABELS["historical_mean_trend_shift"], metric(trip, "historical_mean_trend_shift", "weighted_mae"), COLORS["orange"]),
        (METHOD_LABELS["llm_only_trip_suppression_a1p25"], values["llm_only"], COLORS["purple"]),
        (METHOD_LABELS["zero_shot_llm_rule_tree"], values["zero_rule"], COLORS["blue"]),
        (METHOD_LABELS["llm_rule_small_hist_calibrated_n500"], values["small_rule"], COLORS["orange"]),
        (METHOD_LABELS["global_trip_suppression_a1p25"], metric(trip, "global_trip_suppression_a1p25", "weighted_mae"), COLORS["teal"]),
        (METHOD_LABELS["gated_trip_suppression_a1_d0p15"], values["gated"], COLORS["green"]),
    ]
    draw_bar_chart(slide, chart_values, 0.65, 1.22, 6.25, 3.45, 6.0)
    rows = [
        ["Ordinary XGBoost", f"{values['baseline']:.3f}", f"{values['baseline_bias']:+.3f}", f"{values['baseline_r2']:.3f}"],
        ["CatBoost GPU", f"{values['catboost']:.3f}", f"{values['catboost_bias']:+.3f}", f"{values['catboost_r2']:.3f}"],
        ["LLM-only pressure", f"{values['llm_only']:.3f}", f"{values['llm_only_bias']:+.3f}", f"{values['llm_only_r2']:.3f}"],
        ["Zero-shot tree", f"{values['zero_rule']:.3f}", f"{values['zero_rule_bias']:+.3f}", f"{values['zero_rule_r2']:.3f}"],
        ["Rule + 500 hist.", f"{values['small_rule']:.3f}", f"{values['small_rule_bias']:+.3f}", f"{values['small_rule_r2']:.3f}"],
        ["Hybrid gated", f"{values['gated']:.3f}", f"{values['gated_bias']:+.3f}", f"{values['gated_r2']:.3f}"],
    ]
    small_table(slide, ["Method", "wMAE", "wBias", "wR2"], rows, 7.15, 1.18, [2.05, 0.9, 0.9, 0.8], 0.42, 12, COLORS["teal"])
    metric_card(slide, 7.16, 4.32, 2.05, 0.72, "MAE reduction", f"-{values['mae_reduction']:.1%}", "primary vs XGBoost", COLORS["green"])
    metric_card(slide, 9.45, 4.32, 2.05, 0.72, "Bias reduction", "-99.4%", "absolute weighted bias", COLORS["blue"])
    story_box(
        slide,
        "CatBoost GPU、zero-shot LLM rule tree 和少量历史校准规则都没有超过主方法；加入 gated event adapter 后，wMAE 降到 2.50，bias 接近 0。",
        "The main adapter outperforms stronger tabular, zero-shot rule-tree, and small-data calibration baselines while removing most systematic bias.",
        0.82,
        5.35,
        10.95,
        1.22,
        COLORS["orange"],
    )


def add_statistical_validation_slide(
    prs: Presentation,
    logo: bytes | None,
    ci: pd.DataFrame,
    paired: pd.DataFrame,
) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "09A", "统计验证 / Statistical Validation", "用 household bootstrap 给主结果加不确定性", 11, logo)
    base_mae = ci_values(ci, "traditional_supervised_baseline", "weighted_mae")
    hybrid_mae = ci_values(ci, "gated_llm_correction", "weighted_mae")
    hybrid_bias = ci_values(ci, "gated_llm_correction", "weighted_bias")
    within2 = ci_values(ci, "gated_llm_correction", "weighted_within_2_trips")
    mae_gain = paired_ci_value(paired, "weighted_mae_reduction")
    bias_gain = paired_ci_value(paired, "absolute_bias_reduction")
    rows = [
        ["Baseline wMAE", f"{base_mae[0]:.3f}", f"[{base_mae[1]:.3f}, {base_mae[2]:.3f}]"],
        ["Hybrid wMAE", f"{hybrid_mae[0]:.3f}", f"[{hybrid_mae[1]:.3f}, {hybrid_mae[2]:.3f}]"],
        ["Hybrid bias", f"{hybrid_bias[0]:+.3f}", f"[{hybrid_bias[1]:+.3f}, {hybrid_bias[2]:+.3f}]"],
        ["Within 2 trips", f"{within2[0]:.1%}", f"[{within2[1]:.1%}, {within2[2]:.1%}]"],
    ]
    small_table(slide, ["Metric", "Point", "95% CI"], rows, 0.78, 1.25, [2.7, 1.75, 3.1], 0.58, 12, COLORS["teal"])
    metric_card(slide, 8.35, 1.35, 2.75, 0.92, "MAE gain", f"{mae_gain[0]:.2f}", f"CI [{mae_gain[1]:.2f}, {mae_gain[2]:.2f}]", COLORS["green"])
    metric_card(slide, 8.35, 2.55, 2.75, 0.92, "Bias gain", f"{bias_gain[0]:.2f}", f"CI [{bias_gain[1]:.2f}, {bias_gain[2]:.2f}]", COLORS["blue"])
    story_box(
        slide,
        "Bootstrap 结果说明：主方法的 MAE 改善和 bias 改善不是单个点估计的偶然波动，paired CI 保持为正。",
        "The paired confidence intervals remain positive, supporting the robustness of the event-adaptation gain.",
        0.82,
        5.2,
        10.95,
        1.15,
        COLORS["green"],
    )


def add_tradeoff_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "09B", "误差与偏差权衡 / Error-Bias Tradeoff", "横轴是 |bias|，纵轴是 wMAE，越靠左下越好", 12, logo)
    add_picture(slide, FIGURE_DIR / "mae_bias_tradeoff.png", 0.72, 1.28, 6.75)
    story_box(
        slide,
        "Hybrid 位于低误差、低偏差区域；普通 XGBoost 的偏差较大，说明 2022 的整体高估不能只靠家庭表格变量解决。",
        "The hybrid method moves toward lower error and near-zero bias; the ordinary supervised baseline keeps a clear positive shift.",
        7.75,
        1.42,
        4.1,
        1.58,
        COLORS["green"],
    )
    small_table(
        slide,
        ["读图方式", "含义"],
        [
            ["横轴", "越接近 0，系统性偏差越小"],
            ["纵轴", "越低，家庭出行次数误差越小"],
            ["点大小", "代表 weighted RMSE，反映大误差风险"],
            ["Hybrid", "位于低误差、低偏差区域"],
        ],
        7.75,
        3.12,
        [1.45, 3.1],
        0.38,
        8.3,
        COLORS["teal"],
    )


def add_multi_objective_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "09C", "多目标选择 / Multi-Objective Selection", "从单一排行榜转向规划目标下的 operating point", 13, logo)
    add_picture(slide, PARETO_DIR / "trip_pareto_frontier.png", 0.62, 1.18, 5.9)
    add_picture(slide, PARETO_DIR / "behavior_system_improvement.png", 6.82, 1.2, 5.55)
    points = read_csv_or_empty(PARETO_DIR / "preference_operating_points.csv")
    if not points.empty:
        profile_labels = {
            "calibration_first": "Calibration",
            "balanced_course_report": "Balanced / main",
            "low_cost_deployment": "Low-cost",
        }
        method_labels = {
            "gated_trip_suppression_a1_d0p15": "Gated LLM",
            "global_trip_suppression_a1": "Global prior",
        }
        rows = []
        show_profiles = ["calibration_first", "balanced_course_report", "low_cost_deployment"]
        for row in points[points["profile"].isin(show_profiles)].itertuples(index=False):
            rows.append(
                [
                    profile_labels.get(row.profile, row.profile),
                    method_labels.get(row.method, str(row.method).replace("_", " ")),
                    f"{row.weighted_mae:.3f}",
                    f"{row.abs_weighted_bias:.3f}",
                    f"{int(row.llm_request_cost)}",
                ]
            )
        small_table(
            slide,
            ["Preference", "Selected method", "wMAE", "|bias|", "LLM req."],
            rows,
            0.82,
            5.0,
            [1.65, 2.15, 0.8, 0.8, 0.9],
            0.43,
            12,
            COLORS["teal"],
        )
    story_box(
        slide,
        "关键讲法：LLM 不是直接报一个答案，而是把事件机制转成候选修正；最终由规划目标选择 Pareto operating point。",
        "Key message: LLM supplies event mechanisms; the solver selects a planning operating point.",
        8.05,
        5.18,
        4.1,
        1.0,
        COLORS["orange"],
    )


def add_error_distribution_figure_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "09D", "误差分布与容忍度 / Error Distribution and Tolerance", "误差分布看系统性高估，容忍曲线看 household-level 可用性", 14, logo)
    add_picture(slide, FIGURE_DIR / "error_distribution.png", 0.65, 1.2, 5.85)
    add_picture(slide, FIGURE_DIR / "tolerance_curve.png", 6.85, 1.2, 5.65)
    story_box(
        slide,
        "传统监督模型误差分布整体偏正，说明它系统性高估 2022 出行；事件修正后误差更集中，容忍阈值内覆盖率更高。",
        "The traditional supervised baseline overpredicts many households; event correction shifts errors closer to zero and improves tolerance-level coverage.",
        0.85,
        5.55,
        10.95,
        1.20,
        COLORS["orange"],
    )


def add_event_heterogeneity_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "10A", "情境影响异质性 / Event Heterogeneity", "pressure 分层和先验相关性检验时序机制是否合理", 15, logo)
    add_picture(slide, FIGURE_DIR / "pressure_quintile_gain.png", 0.65, 1.2, 5.85)
    add_picture(slide, FIGURE_DIR / "event_prior_heatmap.png", 7.0, 1.13, 4.95)
    story_box(
        slide,
        "LLM pressure 与目标年机制相符：高 pressure 组需要更强出行折减，remote work、transit avoidance 等先验之间也呈现可解释相关。",
        "The LLM priors behave like event variables, with pressure strata and prior correlations matching target-year mobility mechanisms.",
        0.85,
        5.62,
        10.95,
        1.10,
        COLORS["teal"],
    )


def add_subgroup_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "10C", "分组稳健性 / Equity-Aware Subgroup Check", "比较不同家庭群体中的 MAE 改善与 worst-subgroup error", 17, logo)
    add_picture(slide, FIGURE_DIR / "subgroup_gain_top.png", 0.72, 1.25, 6.6)
    story_box(
        slide,
        "分组结果显示，改进不只来自总体均值；worst-subgroup wMAE 从 8.38 降到 4.71，所有纳入分组均优于历史 XGBoost。",
        "Worst-subgroup wMAE drops from 8.38 to 4.71, and all evaluated subgroups improve over historical XGBoost.",
        7.65,
        1.38,
        4.25,
        1.58,
        COLORS["purple"],
    )
    small_table(
        slide,
        ["汇报时怎么讲", "注意边界"],
        [
            ["预测层面 equity", "worst-subgroup error 下降"],
            ["注意表述边界", "不是因果公平性结论"],
            ["联系规划含义", "识别更难预测的家庭"],
        ],
        7.65,
        3.15,
        [1.75, 2.65],
        0.42,
        8.5,
        COLORS["navy2"],
    )


def add_robustness_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "10B", "稳健性检验 / Robustness Check", "随机置换检验 LLM pressure 是否包含真实排序信息", 16, logo)
    add_picture(slide, ROBUSTNESS_DIR / "permutation_null_mae.png", 0.72, 1.22, 6.2)
    story_box(
        slide,
        "真实 LLM pressure 的误差低于 500 次随机分配；同时 global pressure 也很强，说明主贡献是目标年事件层面的修正，cohort 排序是增量信息。",
        "Real LLM pressure beats random assignment, while the strong global baseline shows that the main signal is event-level correction.",
        7.45,
        1.35,
        4.45,
        1.6,
        COLORS["red"],
    )
    small_table(
        slide,
        ["可能被问到的问题", "我们的回答"],
        [
            ["global rule 是否已经足够？", "它很强；LLM 排序提供增量价值"],
            ["是否用了 2022 标签？", "LLM 输入排除目标值、权重和 ID"],
            ["LLM 能否替代传统模型？", "不能；hybrid 明显更稳"],
            ["是否只是通用下调？", "pre-COVID placebo 中 full suppression 会过度修正"],
        ],
        7.45,
        3.25,
        [2.0, 2.75],
        0.4,
        8.2,
        COLORS["navy2"],
    )


def add_error_insight_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "11", "误差洞察 / Error Insight", "主收益来自消除 target-year over-prediction", 18, logo)
    metric_card(slide, 0.75, 1.28, 2.45, 0.9, "Exact rounded", f"{values['exact']:.1%}", "household accuracy", COLORS["blue"])
    metric_card(slide, 3.48, 1.28, 2.45, 0.9, "Within 2 trips", f"{values['within2']:.1%}", "practical tolerance", COLORS["green"])
    metric_card(slide, 6.21, 1.28, 2.45, 0.9, "Within 3 trips", f"{values['within3']:.1%}", "broad tolerance", COLORS["orange"])
    metric_card(slide, 8.94, 1.28, 2.45, 0.9, "Weighted R2", f"{values['gated_r2']:.3f}", "primary gated rule", COLORS["purple"])
    text_box(slide, "Weighted bias comparison", 0.82, 2.75, 2.3, 0.24, 10.5, COLORS["teal"], True)
    draw_bias_strip(
        slide,
        [
            ("Historical mean", values["mean_bias"], COLORS["gray"]),
            ("Ordinary XGBoost", values["baseline_bias"], COLORS["red"]),
            ("Trend / indicator", values["trend_bias"], COLORS["orange"]),
            ("LLM-only pressure", values["llm_only_bias"], COLORS["purple"]),
            ("Hybrid gated", values["gated_bias"], COLORS["green"]),
        ],
        0.82,
        3.12,
        5.1,
        2.35,
    )
    headers = ["Observation", "Interpretation"]
    rows = [
        ["XGBoost: household signal + positive bias", "history routine is too high for 2022"],
        ["LLM-only: direction useful, calibration weak", "event semantics cannot replace household baseline"],
        ["Hybrid: bias close to zero", "routine demand and event response are separated"],
    ]
    small_table(slide, headers, rows, 6.35, 3.02, [2.75, 3.0], 0.62, 12, COLORS["navy2"])


def add_llm_role_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "12", "LLM 的角色 / What the LLM Adds", "LLM 提供 selector signal 和 correction prior", 19, logo)
    rect(slide, 1.15, 1.28, 10.35, 4.48, COLORS["light"], COLORS["line"])
    rect(slide, 6.3, 1.28, 0.012, 4.48, COLORS["line"])
    rect(slide, 1.15, 3.52, 10.35, 0.012, COLORS["line"])
    text_box(slide, "事件信息弱", 1.45, 1.02, 2.4, 0.22, 8, COLORS["gray"], True)
    text_box(slide, "事件信息强", 8.0, 1.02, 2.4, 0.22, 8, COLORS["gray"], True)
    text_box(slide, "家庭基线弱", 0.18, 2.0, 0.75, 0.46, 8, COLORS["gray"], True, PP_ALIGN.CENTER)
    text_box(slide, "家庭基线强", 0.18, 4.35, 0.75, 0.46, 8, COLORS["gray"], True, PP_ALIGN.CENTER)
    node(slide, 1.75, 1.82, 3.4, 0.8, "Historical mean", "neither household heterogeneity nor event semantics\nwMAE 5.51", COLORS["gray"])
    node(slide, 7.15, 1.82, 3.4, 0.8, "LLM-only pressure", f"event direction without household baseline\nwMAE {values['llm_only']:.2f}", COLORS["purple"])
    node(slide, 1.75, 4.1, 3.4, 0.8, "Ordinary XGBoost", f"household grounding without pandemic semantics\nwMAE {values['baseline']:.2f}", COLORS["red"])
    node(slide, 7.15, 4.1, 3.4, 0.8, "Hybrid gated", f"household baseline + event prior\nwMAE {values['gated']:.2f}", COLORS["green"])
    rect(slide, 1.1, 6.1, 10.45, 0.48, COLORS["pale"], COLORS["line"], True)
    text_box(slide, "汇报讲法", 1.3, 6.2, 1.0, 0.2, 9, COLORS["orange"], True)
    text_box(
        slide,
        f"LLM-only has direction but weaker calibration: wMAE {values['llm_only']:.2f} -> {values['gated']:.2f}; wBias {values['llm_only_bias']:+.2f} -> {values['gated_bias']:+.2f}.",
        2.35,
        6.13,
        8.8,
        0.32,
        10.5,
        COLORS["ink"],
        True,
    )


def add_mode_slide(prs: Presentation, logo: bytes | None, mode: pd.DataFrame, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "13", "综合行为输出 / Behavior Outputs", "从出行次数扩展到方式结构和方式出行量", 20, logo)
    modes = ["private", "walk", "bike", "transit", "taxi", "other"]
    for idx, mode_name in enumerate(modes):
        x = 0.78 + idx * 1.2
        color = COLORS["orange"] if mode_name == "transit" else COLORS["blue"]
        rect(slide, x, 1.42, 0.82, 0.42, color, None, True)
        text_box(slide, mode_name, x + 0.04, 1.53, 0.74, 0.16, 7.2, COLORS["white"], True, PP_ALIGN.CENTER)
    text_box(slide, "household mode-share vector", 2.25, 2.03, 3.25, 0.22, 9, COLORS["gray"], True, PP_ALIGN.CENTER)
    rows = [
        ["Historical XGBoost", f"{values['mode_tv_base']:.4f}", f"{values['transit_base']:.4f}", "routine mode pattern"],
        ["XGB + LLM prior", f"{values['mode_tv_llm']:.4f}", f"{values['transit_llm']:.4f}", "transit correction"],
        ["Improvement", f"{pct_delta(values['mode_tv_base'], values['mode_tv_llm']):.1%}", f"{values['transit_gain']:.1%}", "clearer on transit"],
    ]
    small_table(slide, ["Method", "weighted TV", "transit MAE", "Interpretation"], rows, 0.78, 2.72, [2.25, 1.1, 1.1, 2.15], 0.42, 8.0, COLORS["teal"])
    metric_card(slide, 7.35, 1.38, 2.2, 0.82, "Mode-trip MAE", f"{values['mode_trip_hybrid']:.2f}", "gated count × LLM mode", COLORS["green"])
    metric_card(slide, 9.85, 1.38, 2.2, 0.82, "Reduction", f"{values['mode_trip_gain']:.1%}", "vs traditional combo", COLORS["blue"])
    add_picture(slide, FIGURE_DIR / "mode_component_mae.png", 7.25, 2.55, 4.8)
    story_box(
        slide,
        "方式结构不是主要增益来源，但 transit share 明显改善；这与目标年公共交通规避机制一致。",
        "Mode composition gains are modest overall, but transit-share error improves in the direction suggested by pandemic transit avoidance.",
        0.82,
        5.50,
        10.95,
        1.18,
        COLORS["orange"],
    )


def add_purpose_slide(prs: Presentation, logo: bytes | None, purpose: pd.DataFrame) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "14", "出行目的扩展 / Purpose Composition", "第三个行为维度：为什么出行", 21, logo)
    add_picture(slide, PURPOSE_DIR / "figures" / "purpose_distribution_shift.png", 0.72, 1.18, 5.75)
    if purpose.empty:
        rows = [["missing", "-", "-", "-"]]
    else:
        base = purpose[purpose["method"] == "historical_xgboost"].iloc[0]
        llm = purpose[purpose["method"] == "llm_purpose_prior_a1"].iloc[0]
        global_row = purpose[purpose["method"] == "global_purpose_prior_a1"].iloc[0]
        rows = [
            ["Traditional XGBoost", f"{base.weighted_total_variation:.3f}", f"{base.work_share_weighted_mae:.3f}", "strongest overall"],
            ["LLM purpose prior", f"{llm.weighted_total_variation:.3f}", f"{llm.work_share_weighted_mae:.3f}", "directional components"],
            ["Global purpose prior", f"{global_row.weighted_total_variation:.3f}", f"{global_row.shopping_share_weighted_mae:.3f}", "shopping improves"],
        ]
    small_table(slide, ["Method", "TV", "Component MAE", "Role"], rows, 6.72, 1.28, [2.25, 0.85, 1.35, 1.45], 0.56, 12, COLORS["teal"])
    story_box(
        slide,
        "purpose 结果说明综合目标已经扩展到“为什么出行”；但 LLM prior 目前不是整体最优，适合作为边界和未来工作，而不是主贡献。",
        "Purpose composition broadens the behavior system, while also showing a clear limitation: event priors need stronger task-specific calibration.",
        0.85,
        5.55,
        10.95,
        1.15,
        COLORS["purple"],
    )


def add_method_spectrum_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(
        slide,
        "04",
        "方法谱系 / Method Spectrum",
        "先把 selector branch 和 LLM-corrected XGBoost 串成一条线",
        6,
        logo,
    )
    cards = [
        ("Data-only", "历史数据拟合", values["baseline"], values["baseline_bias"], COLORS["red"]),
        ("Pure LLM prior", "事件方向可用", values["llm_only"], values["llm_only_bias"], COLORS["purple"]),
        ("LLM rule tree", "零样本规则", values["zero_rule"], values["zero_rule_bias"], COLORS["orange"]),
        ("Rule + history", "少量历史校准", values["small_rule"], values["small_rule_bias"], COLORS["teal"]),
        ("Hybrid adapter", "LLM 修正 XGBoost", values["gated"], values["gated_bias"], COLORS["green"]),
    ]
    for idx, (title, note, mae, bias, color) in enumerate(cards):
        x = 0.62 + idx * 2.46
        rect(slide, x, 1.24, 2.15, 1.18, COLORS["light"], COLORS["line"], True)
        rect(slide, x, 1.24, 2.15, 0.06, color)
        text_box(slide, title, x + 0.12, 1.42, 1.92, 0.22, 11.5, color, True, PP_ALIGN.CENTER)
        text_box(slide, note, x + 0.12, 1.73, 1.92, 0.22, 9.2, COLORS["gray"], False, PP_ALIGN.CENTER)
        text_box(slide, f"wMAE {mae:.3f}", x + 0.12, 1.96, 1.92, 0.2, 10.5, COLORS["ink"], True, PP_ALIGN.CENTER)
        text_box(slide, f"bias {bias:+.3f}", x + 0.12, 2.17, 1.92, 0.18, 8.5, COLORS["gray"], False, PP_ALIGN.CENTER)
        if idx < len(cards) - 1:
            arrow(slide, x + 2.19, 1.75, 0.24, 0.14, COLORS["gray"])

    small_table(
        slide,
        ["阶段", "核心假设", "最适合回答的问题", "本项目结论"],
        [
            ["Data-only", "历史样本足够表达未来", "常态年份迁移", "2022 出现系统性高估"],
            ["Pure LLM", "预训练知识可直接预测", "零样本方向判断", "方向有用，数值校准不足"],
            ["LLM rule", "把常识蒸馏为规则树", "老师追问的零样本 baseline", "wMAE 2.602，弱于 hybrid"],
            ["Rule + history", "用少量历史标签校准规则", "data-sparse / cold-start", "合理桥梁，但不是本任务最优"],
            ["Hybrid adapter", "历史模型学 routine，LLM 学 target-year shift", "目标年无标签模型修正", "wMAE 2.502，bias 接近 0"],
        ],
        0.68,
        3.04,
        [1.55, 3.15, 3.15, 3.25],
        0.46,
        12,
        COLORS["navy2"],
    )
    story_box(
        slide,
        "整合口径：模拟预测选择器负责 cold-start / sparse-data 条件下的快速候选预测；LLM-corrected XGBoost 负责有历史标签但目标年情境变化的主任务。二者是同一 temporal-adaptation framework 的两个分支。",
        "The selector branch handles sparse-data simulation, while the correction branch anchors routine mobility in historical survey labels and adapts it with LLM context priors.",
        0.82,
        6.18,
        10.95,
        0.72,
        COLORS["blue"],
    )


def add_scalability_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(
        slide,
        "15",
        "LLM 蒸馏与部署 / Distillation & Deployment",
        "从逐户调用转成 batch prior、规则适配和置信度回退",
        22,
        logo,
    )
    stages = [
        ("Direct LLM", "逐户预测成本高；校准弱"),
        ("Cohort batch", "1,327 cohorts -> 89 prompts"),
        ("Structured priors", "JSON schema -> event variables"),
        ("Rule adapter", "fixed no-label correction"),
        ("Confidence gate", "global fallback; online LLM optional"),
    ]
    for idx, (title, note) in enumerate(stages):
        x = 0.58 + idx * 2.5
        color = [COLORS["red"], COLORS["teal"], COLORS["orange"], COLORS["green"], COLORS["purple"]][idx]
        rect(slide, x, 1.22, 2.18, 1.08, COLORS["light"], COLORS["line"], True)
        rect(slide, x, 1.22, 2.18, 0.055, color)
        text_box(slide, title, x + 0.12, 1.42, 1.94, 0.24, 11.5, color, True, PP_ALIGN.CENTER)
        text_box(slide, note, x + 0.14, 1.73, 1.9, 0.38, 10.0, COLORS["gray"], False, PP_ALIGN.CENTER)
        if idx < len(stages) - 1:
            arrow(slide, x + 2.22, 1.70, 0.24, 0.14, COLORS["gray"])

    metric_card(slide, 0.95, 2.95, 2.35, 0.84, "Household requests", "-98.9%", "7893 -> 89 prompts", COLORS["blue"])
    metric_card(slide, 3.65, 2.95, 2.35, 0.84, "Batch size", "15", "stable JSON parsing", COLORS["teal"])
    metric_card(slide, 6.35, 2.95, 2.35, 0.84, "Main inference", "0 LLM", "adapter is deterministic", COLORS["green"])
    metric_card(slide, 9.05, 2.95, 2.35, 0.84, "Optional fallback", "low-conf", "deployment extension", COLORS["purple"])

    small_table(
        slide,
        ["部署层", "在本研究中的角色", "为什么这样设计"],
        [
            ["Batch prior", "离线生成 cohort 级情境先验", "保留 LLM 泛化能力，避免逐户调用"],
            ["Rule distillation", "把自然语言判断压缩成 s_event 和 fixed adapter", "比直接 LLM 输出 CNTTDHH 更可审计"],
            ["Confidence gate", "低置信 cohort 退回 global pressure", "主实验保持 no-label、可复现"],
            ["LLM fallback", "作为线上部署扩展，不进入主结果", "有预算时处理边界样本"],
        ],
        0.82,
        4.17,
        [2.35, 4.15, 4.5],
        0.42,
        12,
        COLORS["navy2"],
    )
    text_box(
        slide,
        "汇报口径：LLM 不是每户都报答案，而是把目标年 context 蒸馏成结构化 priors 或 selector signals；主实验预测阶段是 deterministic adapter，低置信 LLM 回退属于可部署扩展。",
        0.95,
        6.48,
        11.0,
        0.34,
        12,
        COLORS["ink"],
        True,
    )


def add_contribution_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_frame(slide, "16", "研究逻辑链 / Research Logic", "从问题定义到方法优势的完整闭环", 23, logo)
    small_table(
        slide,
        ["环节", "怎么做", "证据 / 结果"],
        [
            ["问题定义", "2022 是目标年情境变化，不是平滑跨年预测", "普通 XGBoost bias +3.61 trips"],
            ["传统基线", "用历史 NHTS 学 f_hist(X)，保留家庭属性差异", "wMAE 4.34，说明 household signal 有用"],
            ["LLM 进入点", "cohort profile -> s_event，表达 remote work / transit avoidance", "不输入 CNTTDHH、权重和 household ID"],
            ["融合规则", "ŷ = f_hist(X) × correction(s_event)，参数固定", "wMAE 2.50，bias -0.023"],
            ["结果边界", "主增益来自事件层面修正，mode 改善集中在 transit", "不夸大成纯 LLM 或完整 mode-choice model"],
        ],
        0.72,
        1.18,
        [1.65, 5.25, 4.0],
        0.62,
        12,
        COLORS["teal"],
    )
    rect(slide, 0.9, 5.2, 10.7, 0.75, COLORS["pale"], COLORS["line"], True)
    text_box(slide, "一句话故事", 1.12, 5.43, 1.25, 0.24, 9.5, COLORS["orange"], True)
    text_box(slide, "传统模型负责“正常情况下谁会出行更多”，LLM 负责“目标年哪些出行会被压低”，二者相乘得到无标签修正预测。", 2.42, 5.34, 8.8, 0.34, 12, COLORS["ink"], True)


def add_final_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide, COLORS["navy"])
    add_frame(slide, "END", "总结 / Conclusion", "模拟预测选择 + LLM 修正 XGBoost 的时序适应框架", 24, logo, True)
    rect(slide, 0.82, 1.45, 11.15, 2.2, RGBColor(30, 64, 91), RGBColor(71, 85, 105), True)
    text_box(slide, "我们研究的不是“LLM 直接预测出行次数”，而是：\n当目标年情境变化打破历史连续性时，能否用 LLM 的时序泛化能力，为传统 household travel model 提供无标签修正先验。", 1.12, 1.73, 10.5, 0.9, 15, COLORS["white"], True, PP_ALIGN.CENTER)
    metric_card(slide, 1.0, 4.25, 2.45, 0.9, "Accuracy", f"-{values['mae_reduction']:.1%}", "weighted MAE", COLORS["blue"], COLORS["white"])
    metric_card(slide, 3.85, 4.25, 2.45, 0.9, "Bias", f"{values['gated_bias']:+.3f}", "weighted bias", COLORS["green"], COLORS["white"])
    metric_card(slide, 6.7, 4.25, 2.45, 0.9, "Scope", "4 outputs", "trips / modes / purposes", COLORS["orange"], COLORS["white"])
    metric_card(slide, 9.55, 4.25, 2.45, 0.9, "LLM role", "adapter", "not standalone predictor", COLORS["purple"], COLORS["white"])
    text_box(slide, "一句话总结", 1.02, 6.08, 1.35, 0.24, 12, RGBColor(203, 213, 225), True)
    text_box(slide, "用 NHTS 历史数据学习正常出行规律，用 LLM 情境先验泛化目标年变化，并输出家庭层面的出行强度、方式结构和目的结构。", 2.35, 6.0, 8.9, 0.38, 12, COLORS["white"], True)


def add_backup_qa_slide(
    prs: Presentation,
    logo: bytes | None,
    label: str,
    title: str,
    question: str,
    short_answer: str,
    evidence_rows: list[list[str]],
    response: str,
) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_backup_frame(slide, title, question, label, logo)
    rect(slide, 0.72, 1.22, 11.25, 0.82, COLORS["pale"], COLORS["line"], True)
    text_box(slide, "Short answer", 0.95, 1.42, 1.6, 0.2, 10, COLORS["orange"], True)
    text_box(slide, short_answer, 2.45, 1.33, 9.15, 0.36, 13, COLORS["ink"], True)
    small_table(slide, ["Evidence", "Value / implication"], evidence_rows, 0.78, 2.38, [4.4, 6.6], 0.46, 12, COLORS["teal"])
    story_box(slide, response, "", 0.82, 5.55, 10.95, 1.1, COLORS["blue"])


def add_backup_cohort_value_slide(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = blank_slide(prs)
    set_background(slide)
    add_backup_frame(
        slide,
        "Global prior 很强时，cohort LLM 还贡献什么？",
        "Same-alpha decomposition of global event correction and selective cohort refinement",
        "B9",
        logo,
    )
    add_picture(slide, COHORT_PRIOR_DIR / "cohort_prior_value_top_groups.png", 0.72, 1.2, 6.45)
    metric_card(
        slide,
        7.55,
        1.25,
        1.35,
        0.82,
        "Same-alpha gain",
        f"{values['cohort_global_delta']:.4f}",
        "wMAE lower",
        COLORS["green"],
    )
    metric_card(
        slide,
        9.15,
        1.25,
        1.35,
        0.82,
        "Subgroup win",
        f"{values['cohort_subgroup_share']:.1%}",
        "cells vs global",
        COLORS["blue"],
    )
    metric_card(
        slide,
        10.75,
        1.25,
        1.35,
        0.82,
        "Gate active",
        f"{values['cohort_gate_weighted']:.1%}",
        "weighted HHs",
        COLORS["orange"],
    )
    story_box(
        slide,
        "正确解释：global event correction 是主体，cohort-specific LLM pressure 只在差异足够大时启用，提供选择性校准增益。",
        "The safe claim is global event adaptation plus selective cohort refinement, not that cohort ranking explains the full gain.",
        7.55,
        2.35,
        4.55,
        1.35,
        COLORS["teal"],
    )
    small_table(
        slide,
        ["Evidence", "Implication"],
        [
            ["primary vs global-a1", f"{values['cohort_global_delta']:.4f} lower wMAE"],
            ["primary vs global-a1.25", f"{values['cohort_global_a125_delta']:.4f} lower wMAE"],
            ["where it helps most", "lower-income / zero-vehicle / no-worker cells"],
            ["boundary", "some high-vehicle/high-income cells prefer global"],
        ],
        7.55,
        4.05,
        [2.1, 2.45],
        0.42,
        12,
        COLORS["navy2"],
    )


def add_backup_qa_slides(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    add_backup_qa_slide(
        prs,
        logo,
        "B1",
        "为什么不直接让 LLM 构建 2022 决策树？",
        "Why not ask the LLM to build the 2022 decision tree directly?",
        "已经实现为 baseline；方向有用，但数值校准弱于 hybrid adapter。",
        [
            ["Zero-shot LLM rule tree", f"wMAE {values['zero_rule']:.4f}; wBias {values['zero_rule_bias']:+.4f}"],
            ["Pseudo-label tree", f"wMAE {metric(trip, 'zero_shot_pseudo_label_tree', 'weighted_mae'):.4f}; still follows LLM belief labels"],
            ["Primary hybrid gated", f"wMAE {values['gated']:.4f}; wBias {values['gated_bias']:+.4f}"],
            ["Key issue", "No target labels for split thresholds and leaf values"],
        ],
        "答法：LLM 可以生成 qualitative belief tree，但没有 2022 标签时无法学习真实 NHTS 的 split threshold 和 leaf value。我们的设计把数值预测锚定在历史 NHTS household baseline 上，只让 LLM 提供事件机制先验。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B2",
        "有没有使用 2022 标签或 retrospective leakage？",
        "Did the pipeline use 2022 labels or retrospective leakage?",
        "主方法不使用 2022 CNTTDHH 训练或校准；2022 只做最终 evaluation。",
        [
            ["Main method", "gated_trip_suppression_a1_d0p15 fixed before evaluation"],
            ["LLM inputs", "Exclude HOUSEID, CNTTDHH, WTHHFIN"],
            ["Leakage audit", "LLM-facing files pass forbidden-column scan"],
            ["Prospective context", "prospective_event_context_2022 documents event assumptions"],
        ],
        "答法：LLM 输入是 cohort-level household profile 和泛化目标年事件语境，不含目标值、权重和 household ID。我们用 leakage audit、zero-shot baseline、placebo 和 fixed adapter 共同约束 no-label claim。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B3",
        "Global event prior 已经很强，LLM 还贡献什么？",
        "If the global prior is strong, what does the LLM add?",
        "主贡献应表述为 event-level correction；cohort ranking 是增量 refinement。",
        [
            ["Global event prior", "wMAE 2.5223; wBias -0.7477"],
            ["Primary gated adapter", f"wMAE {values['gated']:.4f}; wBias {values['gated_bias']:+.4f}"],
            ["Same-alpha decomposition", f"{values['cohort_global_delta']:.4f} lower wMAE; {values['cohort_subgroup_share']:.1%} subgroup-cell wins"],
            ["Gate usage", f"cohort-specific pressure used for {values['cohort_gate_weighted']:.1%} weighted households"],
            ["Pseudo-event ranked", "best placebo wMAE 2.6274"],
        ],
        "答法：不能夸大成 LLM cohort ranking 独自贡献全部提升。更准确的说法是：global event correction 是主体，LLM cohort prior 在差异足够大的家庭上提供 selective refinement，改善 calibration 和 subgroup robustness。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B4",
        "和 2025-2026 mobility foundation model 有什么区别？",
        "How is this different from recent mobility foundation-model work?",
        "最新工作多是 trajectory / flow / sensor forecasting；我们是 NHTS household survey 的 event adaptation。",
        [
            ["Event-driven LLM mobility", "ELLMob ICLR'26; CausalMob KDD'25"],
            ["Zero-shot / agentic LLM mobility", "AgentMove NAACL'25; AgentMob 2026"],
            ["Efficient LLM mobility pipeline", "ELP-Mob GIS'25"],
            ["Foundation model direction", "UniMob KDD'25; STFM survey'25"],
            ["Our unit", "Household survey record, not trajectory or road node"],
        ],
        "答法：我们借用了 event-generalization 思路，但不声称训练 universal mobility foundation model。LLM 是 structured event-prior generator，数值需求预测仍由 NHTS 历史标签锚定。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B5",
        "Mode / purpose 是否已经解决？",
        "Are mode and purpose choice solved?",
        "没有。trip generation 是主贡献；mode/purpose 是 behavior-system extension。",
        [
            ["Transit-share MAE", "0.0325 -> 0.0269"],
            ["Mode-trip volume MAE", "4.8467 -> 3.2802"],
            ["Purpose composition", "Exploratory; LLM prior not overall best"],
            ["Safe claim", "Planning relevance, not solved full mode-choice model"],
        ],
        "答法：综合目标是预测 household behavior system，但最强结果仍是 target-year trip generation。方式和目的模块说明可扩展到规划相关输出，也明确了未来工作边界。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B6",
        "这些指标怎么用一句话解释？",
        "How should the metrics be interpreted?",
        "wMAE 是加权平均误差，wBias 是系统性高估/低估，越接近 0 越好。",
        [
            ["Historical XGBoost", f"wMAE {values['baseline']:.4f}; wBias {values['baseline_bias']:+.4f}"],
            ["Primary gated", f"wMAE {values['gated']:.4f}; wBias {values['gated_bias']:+.4f}"],
            ["Within 2 trips", f"{values['within2']:.1%} households"],
            ["Within 3 trips", f"{values['within3']:.1%} households"],
        ],
        "答法：MAE 看单个家庭预测误差，bias 看总需求是否系统性偏高或偏低。我们的核心收益不是只降低误差，而是几乎消除了 2022 的系统性高估。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B7",
        "有没有 NHTS 之外的外部验证？",
        "Do we have validation beyond NHTS?",
        "有机制级验证和 PSRC household microdata；但不是 NHTS 2022 的直接复刻。",
        [
            ["ACS work from home", "2019 5.7% -> 2022 15.2%; supports remote-work prior"],
            ["ACS public transit", "2019 5.0% -> 2022 3.1%; supports transit-avoidance prior"],
            ["PSRC microdata", "2017+2019->2023 wMAE 3.4271->3.2498; wBias +1.0532->+0.2605"],
            ["Safe claim", "External mechanism + recovery-transfer support"],
        ],
        "答法：我们用 ACS 证明目标年通勤机制确实发生了远程办公和公共交通下降；用 PSRC 独立 household microdata 做 recovery-transfer 验证；同时用 BTS 做 guardrail，说明 device mobility trips 不能直接当作 NHTS household trip-count 的外部标签。",
    )
    add_backup_qa_slide(
        prs,
        logo,
        "B8",
        "规则蒸馏和低置信 LLM 回退怎么并入本方法？",
        "How do rule distillation and low-confidence LLM fallback fit into the method?",
        "它们是同一条技术路线的部署层：LLM 先验被蒸馏成规则，低置信样本可选回退。",
        [
            ["Distilled rule layer", "LLM priors -> fixed event adapter, not per-household answers"],
            ["Batch prompting", "1,327 cohorts -> 89 prompts with batch size 15"],
            ["Reported experiment", "No target-year labels; low-confidence uses global pressure fallback"],
            ["Deployable extension", "Online LLM fallback only for uncertain cohorts if cost allows"],
        ],
        "答法：可以把“LLM 规则蒸馏 + 置信度回退”理解为我们方法的工程化外壳。论文主结果为了避免 target-year leakage，使用离线 batch 先验和固定 adapter；如果未来做线上系统，可以在低置信 cohort 上再调用 LLM，而不是改变本文主实验口径。",
    )


def create_deck() -> Path:
    logo = extract_template_logo()
    template = TEMPLATE_PATH if TEMPLATE_PATH.exists() else None
    prs = Presentation(template) if template else Presentation()
    if not template:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    clear_template_slides(prs)
    trip, acc, mode, mode_trips, purpose, ci, paired_ci = load_results()
    strong = read_csv_or_empty(PROJECT_ROOT / "outputs" / "strong_baselines" / "strong_tabular_baseline_metrics.csv")

    baseline = metric(trip, "historical_xgboost", "weighted_mae")
    gated = metric(trip, "gated_trip_suppression_a1_d0p15", "weighted_mae")
    llm_only = metric(trip, "llm_only_trip_suppression_a1p25", "weighted_mae")
    transit_base = metric(mode, "historical_xgboost", "transit_share_weighted_mae")
    transit_llm = metric(mode, "llm_transit_avoidance_a1", "transit_share_weighted_mae")
    mode_tv_base = metric(mode, "historical_xgboost", "weighted_total_variation")
    mode_tv_llm = metric(mode, "llm_transit_avoidance_a1", "weighted_total_variation")
    if mode_trips.empty:
        mode_trip_base = float("nan")
        mode_trip_hybrid = float("nan")
    else:
        mode_trip_base = float(
            mode_trips.loc[
                mode_trips["method"] == "traditional_count_x_traditional_mode",
                "weighted_total_mode_trip_mae",
            ].iloc[0]
        )
        mode_trip_hybrid = float(
            mode_trips.loc[
                mode_trips["method"] == "gated_count_x_llm_mode",
                "weighted_total_mode_trip_mae",
            ].iloc[0]
        )
    if strong.empty or "catboost_gpu" not in set(strong["method"]):
        catboost_mae = float("nan")
        catboost_bias = float("nan")
        catboost_r2 = float("nan")
    else:
        catboost_row = strong[strong["method"] == "catboost_gpu"].iloc[0]
        catboost_mae = float(catboost_row["weighted_mae"])
        catboost_bias = float(catboost_row["weighted_bias"])
        catboost_r2 = float(catboost_row["weighted_r2"])
    main_method = "gated_trip_suppression_a1_d0p15"
    values = {
        "baseline": baseline,
        "baseline_bias": metric(trip, "historical_xgboost", "weighted_bias"),
        "baseline_r2": metric(trip, "historical_xgboost", "weighted_r2"),
        "catboost": catboost_mae,
        "catboost_bias": catboost_bias,
        "catboost_r2": catboost_r2,
        "gated": gated,
        "gated_bias": metric(trip, main_method, "weighted_bias"),
        "gated_r2": metric(trip, main_method, "weighted_r2"),
        "llm_only": llm_only,
        "llm_only_bias": metric(trip, "llm_only_trip_suppression_a1p25", "weighted_bias"),
        "llm_only_r2": metric(trip, "llm_only_trip_suppression_a1p25", "weighted_r2"),
        "zero_rule": metric(trip, "zero_shot_llm_rule_tree", "weighted_mae"),
        "zero_rule_bias": metric(trip, "zero_shot_llm_rule_tree", "weighted_bias"),
        "zero_rule_r2": metric(trip, "zero_shot_llm_rule_tree", "weighted_r2"),
        "small_rule": metric(trip, "llm_rule_small_hist_calibrated_n500", "weighted_mae"),
        "small_rule_bias": metric(trip, "llm_rule_small_hist_calibrated_n500", "weighted_bias"),
        "small_rule_r2": metric(trip, "llm_rule_small_hist_calibrated_n500", "weighted_r2"),
        "mean_bias": metric(trip, "historical_mean_only", "weighted_bias"),
        "trend_bias": metric(trip, "historical_mean_trend_shift", "weighted_bias"),
        "mae_reduction": pct_delta(baseline, gated),
        "exact": metric(acc, main_method, "exact_rounded_accuracy"),
        "within2": metric(acc, main_method, "within_2_trips"),
        "within3": metric(acc, main_method, "within_3_trips"),
        "transit_base": transit_base,
        "transit_llm": transit_llm,
        "transit_gain": pct_delta(transit_base, transit_llm),
        "mode_tv_base": mode_tv_base,
        "mode_tv_llm": mode_tv_llm,
        "mode_trip_base": mode_trip_base,
        "mode_trip_hybrid": mode_trip_hybrid,
        "mode_trip_gain": pct_delta(mode_trip_base, mode_trip_hybrid) if mode_trip_base == mode_trip_base else float("nan"),
        "cohort_global_delta": cohort_value("overall_primary_vs_global_a1_mae_delta"),
        "cohort_global_a125_delta": cohort_value("overall_primary_vs_global_a1p25_mae_delta"),
        "cohort_subgroup_share": cohort_value("share_subgroup_cells_primary_beats_global_a1"),
        "cohort_gate_weighted": cohort_value("gated_uses_cohort_prior_share_weighted"),
    }

    add_title_slide(prs, logo, values)
    add_content_slide(prs, logo)
    add_problem_slide(prs, logo, values)
    add_related_work_slide(prs, logo)
    add_data_slide(prs, logo)
    add_method_spectrum_slide(prs, logo, values)
    add_route_slide(prs, logo)
    add_llm_prior_slide(prs, logo)
    add_comparison_slide(prs, logo, trip)
    add_results_slide(prs, logo, trip, values)
    add_statistical_validation_slide(prs, logo, ci, paired_ci)
    add_tradeoff_slide(prs, logo)
    add_multi_objective_slide(prs, logo)
    add_error_distribution_figure_slide(prs, logo)
    add_event_heterogeneity_slide(prs, logo)
    add_robustness_slide(prs, logo)
    add_subgroup_slide(prs, logo)
    add_error_insight_slide(prs, logo, values)
    add_llm_role_slide(prs, logo, values)
    add_mode_slide(prs, logo, mode, values)
    add_purpose_slide(prs, logo, purpose)
    add_scalability_slide(prs, logo)
    add_contribution_slide(prs, logo)
    add_final_slide(prs, logo, values)
    add_backup_qa_slides(prs, logo, values)
    add_backup_cohort_value_slide(prs, logo, values)

    return save_presentation(prs, OUTPUT_PATH)


def save_presentation(prs: Presentation, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(output_path)
        return output_path
    except PermissionError:
        for index in range(2, 20):
            fallback = output_path.with_name(f"{output_path.stem}_v{index}{output_path.suffix}")
            try:
                prs.save(fallback)
                LOGGER.warning("Primary PPT is locked; wrote fallback file: %s", fallback)
                return fallback
            except PermissionError:
                continue
        raise


def main() -> None:
    configure_logging()
    path = create_deck()
    LOGGER.info("Wrote template-based presentation: %s", path)


if __name__ == "__main__":
    main()
