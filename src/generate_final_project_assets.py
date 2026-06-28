"""Generate final report, figures, and presentation assets."""

from __future__ import annotations

import csv
import logging
import sys
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / "src"))

from run_household_baseline import TARGET_COLUMN, WEIGHT_COLUMN  # noqa: E402
from run_label_free_llm_adaptation import (  # noqa: E402
    add_event_pressure,
    apply_multiplicative_pressure,
    load_2022_with_llm_features,
    make_historical_mean_prediction,
    make_delta_gated_pressure,
    make_global_pressure,
)
from run_llm_residual_adaptation import resolve_device, train_history_model  # noqa: E402


LOGGER = logging.getLogger(__name__)
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "final_project"
FIGURE_DIR = OUTPUT_DIR / "figures"
LABEL_FREE_DIR = PROJECT_ROOT / "outputs" / "label_free_llm_adaptation"
ZERO_SHOT_RULE_TREE_PATH = (
    PROJECT_ROOT
    / "outputs"
    / "zero_shot_llm_rule_tree_baseline"
    / "zero_shot_llm_rule_tree_metrics.csv"
)
IRRELEVANT_PSEUDO_EVENT_PATH = (
    PROJECT_ROOT
    / "outputs"
    / "irrelevant_pseudo_event_placebo"
    / "irrelevant_pseudo_event_placebo_metrics.csv"
)
SMALL_DATA_CALIBRATION_PATH = (
    PROJECT_ROOT
    / "outputs"
    / "llm_rule_small_data_calibration"
    / "method_spectrum_metrics.csv"
)

SELECTED_METHODS = [
    "historical_mean_only",
    "historical_xgboost",
    "historical_mean_trend_shift",
    "llm_only_trip_suppression_a1p25",
    "global_trip_suppression_a1p25",
    "random_trip_suppression_a1p25",
    "gated_trip_suppression_a1_d0p15",
]
PLOT_LABELS = {
    "historical_mean_only": "Historical mean",
    "historical_xgboost": "Traditional supervised baseline",
    "historical_mean_trend_shift": "Historical trend shift",
    "llm_only_trip_suppression_a1p25": "LLM-only pressure",
    "global_trip_suppression_a1p25": "Global event prior",
    "random_trip_suppression_a1p25": "Random prior control",
    "gated_trip_suppression_a1_d0p15": "Gated LLM correction",
}
COLORS = {
    "navy": "#17324D",
    "blue": "#3B82F6",
    "cyan": "#06B6D4",
    "green": "#16A34A",
    "orange": "#F97316",
    "red": "#DC2626",
    "gray": "#64748B",
    "light": "#F8FAFC",
    "dark": "#0F172A",
}


@dataclass(frozen=True)
class GeneratedAssets:
    """Paths to final generated artifacts."""

    report_path: Path
    storyboard_path: Path
    pptx_path: Path
    metrics_path: Path
    accuracy_path: Path


def configure_logging() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")


def ensure_dirs() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)


def load_label_free_metrics() -> pd.DataFrame:
    path = LABEL_FREE_DIR / "label_free_llm_adaptation_metrics.csv"
    metrics = pd.read_csv(path)
    missing = set(SELECTED_METHODS).difference(metrics["method"].unique())
    if missing:
        raise ValueError(f"Missing selected methods in {path}: {sorted(missing)}")
    return metrics


def load_optional_metric_row(path: Path, method: str) -> pd.Series | None:
    if not path.exists():
        return None
    metrics = pd.read_csv(path)
    rows = metrics.loc[metrics["method"] == method]
    if rows.empty:
        return None
    return rows.iloc[0]


def aggregate_metrics(metrics: pd.DataFrame) -> pd.DataFrame:
    selected = metrics[metrics["method"].isin(SELECTED_METHODS)].copy()
    summary = (
        selected.groupby("method")
        .agg(
            runs=("method", "count"),
            mae=("mae", "mean"),
            rmse=("rmse", "mean"),
            bias=("bias", "mean"),
            r2=("r2", "mean"),
            weighted_mae=("weighted_mae", "mean"),
            weighted_rmse=("weighted_rmse", "mean"),
            weighted_bias=("weighted_bias", "mean"),
            weighted_r2=("weighted_r2", "mean"),
            weighted_target_mean=("weighted_target_mean", "mean"),
            weighted_prediction_mean=("weighted_prediction_mean", "mean"),
        )
        .reset_index()
    )
    order = {method: index for index, method in enumerate(SELECTED_METHODS)}
    summary["method_order"] = summary["method"].map(order)
    summary = summary.sort_values("method_order").drop(columns=["method_order"])

    baseline = summary[summary["method"] == "historical_xgboost"].iloc[0]
    summary["weighted_mae_reduction_pct"] = (
        100.0 * (baseline["weighted_mae"] - summary["weighted_mae"]) / baseline["weighted_mae"]
    )
    summary["weighted_rmse_reduction_pct"] = (
        100.0 * (baseline["weighted_rmse"] - summary["weighted_rmse"]) / baseline["weighted_rmse"]
    )
    summary["abs_weighted_bias_reduction_pct"] = (
        100.0
        * (abs(baseline["weighted_bias"]) - summary["weighted_bias"].abs())
        / abs(baseline["weighted_bias"])
    )
    return summary


def weighted_average(values: np.ndarray, weights: np.ndarray) -> float:
    return float(np.average(values, weights=weights))


def load_2022_predictions() -> tuple[pd.DataFrame, dict[str, np.ndarray]]:
    device = resolve_device("cuda")
    full_frame = pd.read_csv(PROJECT_ROOT / "data" / "processed" / "household_harmonized.csv", low_memory=False)
    model, feature_columns = train_history_model(full_frame, 200, device)
    frame = load_2022_with_llm_features(
        PROJECT_ROOT / "data" / "processed" / "household_harmonized.csv",
        PROJECT_ROOT / "outputs" / "llm_event_features" / "household_cohort_profiles.csv",
        PROJECT_ROOT
        / "outputs"
        / "llm_event_features"
        / "cursor_api_full_gpt55_low_c15"
        / "validated"
        / "llm_event_features_normalized.csv",
    )
    frame["base_prediction"] = np.maximum(model.predict(frame[feature_columns]), 0.0)
    frame = add_event_pressure(frame)

    base_prediction = frame["base_prediction"].to_numpy(dtype=float)
    weights = frame[WEIGHT_COLUMN].to_numpy(dtype=float)
    trip_pressure = frame["llm_trip_suppression_pressure"].to_numpy(dtype=float)
    global_pressure = make_global_pressure(trip_pressure, weights)
    gated_pressure = make_delta_gated_pressure(trip_pressure, global_pressure, 0.15)
    historical_mean_prediction = make_historical_mean_prediction(full_frame, len(frame))

    predictions = {
        "historical_xgboost": base_prediction,
        "llm_only_trip_suppression_a1p25": apply_multiplicative_pressure(
            historical_mean_prediction, trip_pressure, 1.25, 0.05
        ),
        "llm_trip_suppression_a1p25": apply_multiplicative_pressure(
            base_prediction, trip_pressure, 1.25, 0.05
        ),
        "gated_trip_suppression_a1_d0p15": apply_multiplicative_pressure(
            base_prediction, gated_pressure, 1.0, 0.05
        ),
    }
    return frame, predictions


def build_household_accuracy(frame: pd.DataFrame, predictions: dict[str, np.ndarray]) -> pd.DataFrame:
    y_true = frame[TARGET_COLUMN].to_numpy(dtype=float)
    weights = frame[WEIGHT_COLUMN].to_numpy(dtype=float)
    rows: list[dict[str, str | int | float]] = []
    for method, predicted in predictions.items():
        abs_error = np.abs(predicted - y_true)
        error = predicted - y_true
        rounded = np.rint(predicted)
        rows.append(
            {
                "method": method,
                "rows": len(y_true),
                "unweighted_mae": float(np.mean(abs_error)),
                "unweighted_rmse": float(np.sqrt(np.mean(error**2))),
                "unweighted_bias": float(np.mean(error)),
                "weighted_mae": weighted_average(abs_error, weights),
                "weighted_rmse": float(np.sqrt(weighted_average(error**2, weights))),
                "weighted_bias": weighted_average(error, weights),
                "exact_rounded_accuracy": float(np.mean(rounded == y_true)),
                "within_1_trip": float(np.mean(abs_error <= 1.0)),
                "within_2_trips": float(np.mean(abs_error <= 2.0)),
                "within_3_trips": float(np.mean(abs_error <= 3.0)),
                "weighted_within_1_trip": weighted_average((abs_error <= 1.0).astype(float), weights),
                "weighted_within_2_trips": weighted_average((abs_error <= 2.0).astype(float), weights),
                "weighted_within_3_trips": weighted_average((abs_error <= 3.0).astype(float), weights),
            }
        )
    return pd.DataFrame(rows)


def save_csv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, quoting=csv.QUOTE_MINIMAL)


def set_plot_style() -> None:
    plt.rcParams.update(
        {
            "figure.dpi": 150,
            "savefig.dpi": 220,
            "font.size": 11,
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "xtick.labelsize": 9,
            "ytick.labelsize": 10,
            "font.family": "DejaVu Sans",
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )


def save_metric_comparison(summary: pd.DataFrame) -> Path:
    path = FIGURE_DIR / "metric_comparison.png"
    plot_frame = summary[summary["method"].isin(SELECTED_METHODS)].copy()
    labels = [PLOT_LABELS[method] for method in plot_frame["method"]]
    x = np.arange(len(plot_frame))
    width = 0.36

    fig, ax = plt.subplots(figsize=(11, 5.5))
    ax.bar(x - width / 2, plot_frame["weighted_mae"], width, color=COLORS["blue"], label="Weighted MAE")
    ax.bar(x + width / 2, plot_frame["weighted_rmse"], width, color=COLORS["cyan"], label="Weighted RMSE")
    ax.set_title("Prediction Error on 2022 Households")
    ax.set_ylabel("Trips")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=22, ha="right")
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.22)
    for container in ax.containers:
        ax.bar_label(container, fmt="%.2f", padding=2, fontsize=8)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_bias_r2_tradeoff(summary: pd.DataFrame) -> Path:
    path = FIGURE_DIR / "bias_r2_tradeoff.png"
    methods = ["historical_xgboost", "llm_only_trip_suppression_a1p25", "gated_trip_suppression_a1_d0p15"]
    plot_frame = summary[summary["method"].isin(methods)].set_index("method").loc[methods].reset_index()
    labels = [PLOT_LABELS[method] for method in plot_frame["method"]]
    x = np.arange(len(plot_frame))

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5))
    axes[0].bar(x, plot_frame["weighted_bias"], color=[COLORS["red"], COLORS["orange"], COLORS["green"]])
    axes[0].axhline(0, color=COLORS["dark"], linewidth=1)
    axes[0].set_title("Systematic Bias")
    axes[0].set_ylabel("Weighted bias")
    axes[0].set_xticks(x)
    axes[0].set_xticklabels(labels, rotation=20, ha="right")
    axes[0].grid(axis="y", alpha=0.22)

    axes[1].bar(x, plot_frame["weighted_r2"], color=[COLORS["gray"], COLORS["blue"], COLORS["green"]])
    axes[1].axhline(0, color=COLORS["dark"], linewidth=1)
    axes[1].set_title("Weighted R2")
    axes[1].set_ylabel("Weighted R2")
    axes[1].set_xticks(x)
    axes[1].set_xticklabels(labels, rotation=20, ha="right")
    axes[1].grid(axis="y", alpha=0.22)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_household_accuracy(accuracy: pd.DataFrame) -> Path:
    path = FIGURE_DIR / "household_tolerance_accuracy.png"
    methods = ["historical_xgboost", "llm_only_trip_suppression_a1p25", "gated_trip_suppression_a1_d0p15"]
    plot_frame = accuracy.set_index("method").loc[methods].reset_index()
    labels = [PLOT_LABELS[method] for method in plot_frame["method"]]
    metrics = ["exact_rounded_accuracy", "within_1_trip", "within_2_trips", "within_3_trips"]
    metric_labels = ["Exact", "Within 1 trip", "Within 2 trips", "Within 3 trips"]
    x = np.arange(len(metrics))
    width = 0.24

    fig, ax = plt.subplots(figsize=(10.5, 5.2))
    for index, (_, row) in enumerate(plot_frame.iterrows()):
        values = [100.0 * row[column] for column in metrics]
        ax.bar(x + (index - 1) * width, values, width, label=labels[index])
    ax.set_title("Household-Level Hit Rates")
    ax.set_ylabel("Households (%)")
    ax.set_xticks(x)
    ax.set_xticklabels(metric_labels)
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.22)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_pressure_distribution(frame: pd.DataFrame) -> Path:
    path = FIGURE_DIR / "event_pressure_distribution.png"
    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.hist(frame["trip_suppression_risk"], bins=28, color=COLORS["blue"], alpha=0.75, label="Trip suppression risk")
    ax.hist(
        frame["llm_recovery_adjusted_pressure"],
        bins=28,
        color=COLORS["orange"],
        alpha=0.55,
        label="Recovery-adjusted pressure",
    )
    ax.set_title("LLM Event-Prior Distribution Across 2022 Households")
    ax.set_xlabel("Score")
    ax.set_ylabel("Household count")
    ax.legend(frameon=False)
    ax.grid(axis="y", alpha=0.22)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_subgroup_gains() -> Path:
    path = FIGURE_DIR / "subgroup_llm_gains.png"
    comparison = pd.read_csv(LABEL_FREE_DIR / "label_free_subgroup_comparison.csv")
    top = comparison[comparison["rows"] >= 50].sort_values("llm_vs_global_mae_delta", ascending=False).head(8)
    labels = top["subgroup_column"] + "=" + top["subgroup_value"].astype(str)

    fig, ax = plt.subplots(figsize=(9.5, 4.8))
    ax.barh(labels[::-1], top["llm_vs_global_mae_delta"].iloc[::-1], color=COLORS["green"])
    ax.axvline(0, color=COLORS["dark"], linewidth=1)
    ax.set_title("Where LLM Cohort Ranking Beats Global Downscaling")
    ax.set_xlabel("MAE gain vs global control")
    ax.grid(axis="x", alpha=0.22)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_workflow_diagram() -> Path:
    path = FIGURE_DIR / "method_workflow.png"
    fig, ax = plt.subplots(figsize=(12, 5))
    ax.axis("off")
    boxes = [
        ("Historical NHTS\n2001/2009/2017", 0.05, 0.58, COLORS["gray"]),
        ("Routine mobility\npredictor", 0.28, 0.58, COLORS["blue"]),
        ("2022 household\ncohort profile", 0.05, 0.18, COLORS["gray"]),
        ("LLM event priors\n(no target labels)", 0.28, 0.18, COLORS["orange"]),
        ("Label-free\ncorrection rule", 0.55, 0.38, COLORS["green"]),
        ("2022 household\ntrip-count prediction", 0.78, 0.38, COLORS["navy"]),
    ]
    for text, x, y, color in boxes:
        patch = plt.Rectangle((x, y), 0.17, 0.22, facecolor=color, alpha=0.94, transform=ax.transAxes)
        ax.add_patch(patch)
        ax.text(x + 0.085, y + 0.11, text, color="white", ha="center", va="center", fontsize=12, weight="bold")
    arrows = [((0.22, 0.69), (0.28, 0.69)), ((0.22, 0.29), (0.28, 0.29)), ((0.45, 0.69), (0.55, 0.49)), ((0.45, 0.29), (0.55, 0.49)), ((0.72, 0.49), (0.78, 0.49))]
    for start, end in arrows:
        ax.annotate("", xy=end, xytext=start, xycoords="axes fraction", arrowprops={"arrowstyle": "->", "lw": 2.0})
    ax.text(
        0.5,
        0.04,
        "2022 CNTTDHH is used only for final evaluation, not for training or calibration.",
        ha="center",
        va="center",
        fontsize=11,
        color=COLORS["dark"],
    )
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def pct(value: float) -> str:
    return f"{100.0 * value:.1f}%"


def write_report(summary: pd.DataFrame, accuracy: pd.DataFrame, figure_paths: dict[str, Path]) -> Path:
    path = OUTPUT_DIR / "final_project_report.md"
    baseline = summary[summary["method"] == "historical_xgboost"].iloc[0]
    gated = summary[summary["method"] == "gated_trip_suppression_a1_d0p15"].iloc[0]
    gated_acc = accuracy[accuracy["method"] == "gated_trip_suppression_a1_d0p15"].iloc[0]
    zero_rule = load_optional_metric_row(ZERO_SHOT_RULE_TREE_PATH, "zero_shot_llm_rule_tree")
    zero_tree = load_optional_metric_row(ZERO_SHOT_RULE_TREE_PATH, "zero_shot_pseudo_label_tree")
    small_500 = load_optional_metric_row(SMALL_DATA_CALIBRATION_PATH, "llm_rule_small_hist_calibrated_n500")
    small_full = load_optional_metric_row(SMALL_DATA_CALIBRATION_PATH, "llm_rule_full_history_calibrated")
    pseudo_ranked = load_optional_metric_row(IRRELEVANT_PSEUDO_EVENT_PATH, "pseudo_random_hash_rank_matched_a1")
    pseudo_gated = load_optional_metric_row(
        IRRELEVANT_PSEUDO_EVENT_PATH,
        "pseudo_random_hash_rank_matched_gated_a1_d0p15",
    )

    lines = [
        "# Final Project Report",
        "",
        "## Core Claim",
        "",
        "This project formulates 2022 NHTS household travel prediction as event-driven temporal adaptation. "
        "Historical models learn routine mobility, but the post-pandemic wave contains event mechanisms such as "
        "remote work substitution, transit avoidance, online delivery substitution, and uneven recovery. "
        "The main contribution is a label-free hybrid adapter: LLM-derived event priors correct a historical "
        "routine-mobility predictor, while 2022 trip-count labels are used only for final evaluation.",
        "",
        "## Research Questions",
        "",
        "- RQ1: How severely do historical household travel models overpredict 2022 post-pandemic trip generation?",
        "- RQ2: Can event priors reduce this bias without using 2022 `CNTTDHH` labels for training or calibration?",
        "- RQ3: Does the LLM replace historical prediction, or is the stronger design a hybrid of routine mobility and event semantics?",
        "",
        "## Literature Position",
        "",
        "The 2025-2026 literature grounding is summarized in `plan/literature_grounding_2026.md`. "
        "The closest directions are event-driven LLM mobility generation (ELLMob, ICLR 2026), "
        "LLM-derived causal public-event features for mobility prediction (CausalMob, KDD 2025), "
        "zero-shot LLM mobility agents (AgentMove, NAACL 2025), efficient/evidence-grounded LLM mobility agents "
        "(AgentMob, 2026 preprint), efficient LLM mobility pipelines "
        "(ELP-Mob, SIGSPATIAL/GIS 2025), and universal mobility foundation models (UniMob, KDD 2025).",
        "",
        "Our gap is survey-based household mobility under a post-pandemic event shift: most recent work targets "
        "trajectories, flows, traffic sensors, or public-event time series, while this project uses historical NHTS "
        "labels to ground routine household demand and uses LLM priors only for no-label event adaptation.",
        "",
        "## Main Result",
        "",
        "| Method | Weighted MAE | Weighted RMSE | Weighted Bias | Weighted R2 |",
        "|---|---:|---:|---:|---:|",
    ]
    for method in SELECTED_METHODS:
        row = summary[summary["method"] == method].iloc[0]
        lines.append(
            f"| {PLOT_LABELS[method]} | {row.weighted_mae:.4f} | {row.weighted_rmse:.4f} | "
            f"{row.weighted_bias:.4f} | {row.weighted_r2:.4f} |"
        )
    for label, row in [
        ("Zero-shot LLM rule tree", zero_rule),
        ("Zero-shot pseudo-label tree", zero_tree),
        ("LLM rule + 500-history calibration", small_500),
    ]:
        if row is not None:
            lines.append(
                f"| {label} | {row.weighted_mae:.4f} | {row.weighted_rmse:.4f} | "
                f"{row.weighted_bias:.4f} | {row.weighted_r2:.4f} |"
            )
    lines.extend(
        [
            "",
            "## Temporal Transfer Validation",
            "",
            "Before interpreting 2022 as an event-shift target, we checked routine cross-year transfer. "
            "This table is a diagnostic transfer check, not the final ordinary-XGBoost comparison row reported above:",
            "",
            "| Check | Weighted MAE | Weighted Bias | R2 |",
            "|---|---:|---:|---:|",
            "| 2001 -> 2009 | 3.5207 | +0.5818 | 0.4413 |",
            "| 2001+2009 -> 2017 | 4.1272 | +1.2230 | 0.2505 |",
            "| 2001+2009+2017 -> 2022 | 4.4062 | +3.7202 | -0.5055 |",
            "",
            "The pre-COVID checks have mean absolute weighted bias `0.9024`, while the diagnostic 2022 transfer check has absolute weighted bias `3.7202`. The final ordinary-XGBoost comparison row in the main table has weighted bias `+3.6052`. Both estimates support the same problem framing: 2022 is a stronger post-pandemic event shift rather than an ordinary transfer year.",
            "",
            "## External Mechanism Validation",
            "",
            "We added an external validation and compatibility audit in `outputs/external_validation/`. The external evidence is deliberately scoped as mechanism-level validation rather than household-level MAE.",
            "",
            "| External source | Finding | Relevance |",
            "|---|---|---|",
            "| ACS commuting brief | Worked-from-home commute share remains much higher in 2022 than 2019: `5.7% -> 15.2%`. | Supports `remote_work_substitution` event prior. |",
            "| ACS commuting brief | Public-transportation commute share remains lower in 2022 than 2019: `5.0% -> 3.1%`. | Supports `transit_avoidance` event prior. |",
            "| BTS/UMD daily mobility | BTS trips/person changes only `-1.95%` from 2019 to 2022, while NHTS diary `CNTTDHH` has a much larger 2017-to-2022 shift. | Shows BTS device mobility is not a direct numeric label for NHTS household trips. |",
            "| PSRC household travel survey | 2017+2019->2023 household/day pre/post replication wMAE changes `3.4271 -> 3.2498`; wBias changes `+1.0532 -> +0.2605` using an ACS-derived remote-work suppression factor. | Provides household-level external microdata evidence that event-scale suppression improves transfer on an independent travel survey. |",
            "",
            "Interpretation: external data supports the event semantics used by the LLM adapter and includes a household-level external pre/post replication. The PSRC result is not direct numerical validation of NHTS 2022 household predictions because PSRC is a regional survey with different sampling and diary protocols, but it strengthens the paper claim that fixed event priors can improve label-free transfer under post-pandemic survey shifts. The external adapter is also no-label: it is specified from ACS event context, without target-year PSRC label calibration.",
            "",
            "## Improvement Over Historical Predictor",
            "",
            f"- Primary strict no-label row: `{gated.method}`.",
            f"- Primary weighted MAE drops from `{baseline.weighted_mae:.4f}` to `{gated.weighted_mae:.4f}` "
            f"({gated.weighted_mae_reduction_pct:.2f}% reduction).",
            f"- Primary weighted RMSE drops from `{baseline.weighted_rmse:.4f}` to `{gated.weighted_rmse:.4f}` "
            f"({gated.weighted_rmse_reduction_pct:.2f}% reduction).",
            f"- Primary absolute weighted bias drops by `{gated.abs_weighted_bias_reduction_pct:.2f}%`.",
            "- LLM-only pressure improves over naive historical baselines but remains weaker than the hybrid adapter, "
            "supporting the design choice that LLMs provide event semantics rather than standalone household predictions.",
            (
                f"- Zero-shot LLM rule tree reaches weighted MAE `{zero_rule.weighted_mae:.4f}`, "
                f"and the pseudo-label tree distilled from it reaches `{zero_tree.weighted_mae:.4f}`; "
                "both are weaker and more biased than the primary hybrid adapter."
                if zero_rule is not None and zero_tree is not None
                else ""
            ),
            (
                f"- LLM rule + 500 historical calibration reaches weighted MAE `{small_500.weighted_mae:.4f}`; "
                "it is a useful bridge baseline for data-sparse settings, but it reintroduces positive 2022 bias."
                if small_500 is not None
                else ""
            ),
            "",
            "## Household-Level Accuracy",
            "",
            f"- Gated MAE: `{gated_acc.unweighted_mae:.4f}` trips per household.",
            f"- Gated RMSE: `{gated_acc.unweighted_rmse:.4f}` trips per household.",
            f"- Exact rounded hit rate: `{pct(gated_acc.exact_rounded_accuracy)}`.",
            f"- Within 1 trip: `{pct(gated_acc.within_1_trip)}`.",
            f"- Within 2 trips: `{pct(gated_acc.within_2_trips)}`.",
            f"- Within 3 trips: `{pct(gated_acc.within_3_trips)}`.",
            "",
            "## Statistical Validation",
            "",
            "Household bootstrap resampling gives the following 95% confidence intervals:",
            "",
            "- Traditional supervised baseline weighted MAE: `[4.2418, 4.4312]`.",
            "- Gated LLM correction weighted MAE: `[2.4292, 2.5808]`.",
            "- Paired weighted-MAE reduction: `[1.7422, 1.9267]`.",
            "- Paired absolute-bias reduction: `[3.3582, 3.6696]`.",
            "",
            "This supports the claim that the main trip-count improvement is not a single point-estimate artifact.",
            "",
            "## Behavior-System Extension",
            "",
            "The project now reports household travel behavior as a multi-output system:",
            "",
            "1. Trip generation: household `CNTTDHH`.",
            "2. Mode composition: household mode-share vector from `TRPTRANS`.",
            "3. Mode-specific trip volume: predicted total trips multiplied by predicted mode shares.",
            "4. Purpose composition: household purpose-share vector from `TRIPPURP`.",
            "",
            "Mode-specific trip-volume results:",
            "",
            "- Traditional count x traditional mode total mode-trip MAE: `4.8467`.",
            "- Gated count x LLM mode total mode-trip MAE: `3.2802`.",
            "",
            "Purpose-composition results:",
            "",
            "- Traditional XGBoost purpose weighted TV: `0.5661`.",
            "- LLM purpose prior is not the overall best row, but it adds a third behavior dimension and exposes a clear future-work target for purpose-specific event adaptation.",
            "",
            "## Efficiency",
            "",
            "- Household rows in 2022: `7,893`.",
            "- LLM cohort prompts: `1,327`.",
            "- LLM request reduction: `83.2%`, or about `5.95x` fewer requests than household-level prompting.",
            "- Batch prompting with batch size 15 further compresses `1,327` cohort prompts into `89` batch prompts, a `93.29%` request reduction relative to one-cohort prompts.",
            "",
            "## LLM Generalization Role",
            "",
            "The LLM should be framed as an event-generalization module, not as a direct predictor. It maps pandemic mechanisms such as remote work, transit avoidance, online delivery substitution, and uneven recovery onto unlabeled household cohorts. A prospective event-context file is included at `plan/prospective_event_context_2022.md` to make this role more auditable and reduce retrospective leakage risk.",
            "",
            "## Why Not a Zero-Shot LLM Decision Tree",
            "",
            "A decision tree needs labels to learn split thresholds and leaf-level numerical predictions. Without 2022 `CNTTDHH` labels, an LLM-generated tree would be a belief tree or a synthetic-label model rather than a data-fitted 2022 tree. This is why the project uses the LLM as an event-prior generator and keeps numerical prediction grounded in a historical household model trained on real NHTS data.",
            (
                f"We now include this as an explicit ablation. A zero-shot LLM-style semantic rule tree reaches weighted MAE `{zero_rule.weighted_mae:.4f}`, while a pseudo-label tree reaches `{zero_tree.weighted_mae:.4f}`. The primary hybrid adapter remains better at weighted MAE `{gated.weighted_mae:.4f}`."
                if zero_rule is not None and zero_tree is not None
                else ""
            ),
            "",
            "## LLM Rules Plus Small Historical Calibration",
            "",
            (
                f"We also include the collaborator-proposed bridge route: let an LLM-style rule structure define routine demand leaves, calibrate leaf values with historical samples, and then apply the same 2022 event factor. With 500 historical calibration rows, weighted MAE is `{small_500.weighted_mae:.4f}`; with full-history rule calibration, weighted MAE is `{small_full.weighted_mae:.4f}`. This supports the method spectrum but also shows why routine historical calibration alone can overpredict under a post-pandemic shift."
                if small_500 is not None and small_full is not None
                else ""
            ),
            "",
            "## Robustness Check",
            "",
            "We ran additional robustness checks in `outputs/robustness_checks/`.",
            "",
            "- 500-run permutation control for the primary gated rule: actual weighted MAE `2.5023`, random-permutation mean `2.5808`, empirical p-value `0.0020`.",
            "- Same-alpha global pressure remains strong: primary gated weighted MAE `2.5023` vs global-a1 weighted MAE `2.5531`.",
            (
                f"- LLM rule + small historical calibration is a coherent bridge baseline but not a replacement: 500-row calibration weighted MAE `{small_500.weighted_mae:.4f}`."
                if small_500 is not None
                else ""
            ),
            (
                f"- Irrelevant pseudo-event placebo controls are weaker than the primary method: best ranked pseudo-event weighted MAE `{pseudo_ranked.weighted_mae:.4f}`, best gated pseudo-event weighted MAE `{pseudo_gated.weighted_mae:.4f}`."
                if pseudo_ranked is not None and pseudo_gated is not None
                else ""
            ),
            "- Leakage scan passes for LLM-facing profile/feature files: they exclude `HOUSEID`, `CNTTDHH`, and `WTHHFIN`.",
            "",
            "Interpretation for the course report: the dominant contribution is event-level label-free adaptation. Cohort-specific LLM ranking provides measurable incremental signal, but it should not be described as the sole source of improvement.",
            "",
            "## Figures",
            "",
        ]
    )
    for name, figure_path in figure_paths.items():
        rel_path = figure_path.relative_to(OUTPUT_DIR).as_posix()
        lines.append(f"- `{name}`: `{rel_path}`")
    lines.extend(
        [
            "",
            "## Presentation Framing",
            "",
            "Avoid framing the project as a generic feature-only forecasting improvement. The cleaner narrative is: "
            "the traditional supervised baseline fails under a rare event; LLMs provide event semantics "
            "that can be distilled into a lightweight correction rule.",
            "",
            "Do not overstate the mode-composition or purpose-composition extensions. The strongest result remains trip generation under event-driven temporal adaptation; the extensions show a broader behavior system and planning relevance.",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_storyboard() -> Path:
    path = PROJECT_ROOT / "plan" / "ppt_storyboard.md"
    lines = [
        "# PPT Storyboard",
        "",
        "## Slide 1: Title",
        "- Label-Free LLM Event Adaptation for Post-Pandemic Household Travel Prediction",
        "- Key phrase: no 2022 labels for training.",
        "",
        "## Slide 2: Problem",
        "- Historical routine mobility assumptions break in 2022.",
        "- COVID changed commuting, activity participation, delivery substitution, and recovery behavior.",
        "",
        "## Slide 3: Research Question",
        "- Can LLM event priors repair 2022 prediction without target-year labels?",
        "",
        "## Slide 4: Method",
        "- Traditional supervised baseline learns routine mobility.",
        "- LLM produces cohort-level pandemic priors.",
        "- A fixed correction rule distills the priors into prediction adjustment.",
        "",
        "## Slide 5: LLM Priors",
        "- trip suppression, remote work, transit avoidance, online delivery, recovery sensitivity.",
        "- 1,327 cohorts cover 7,893 households.",
        "",
        "## Slide 6: Main Result",
        "- Primary gated weighted MAE reduction: 42.31%.",
        "- Primary gated weighted RMSE reduction: 32.22%.",
        "- Primary gated absolute weighted-bias reduction: 99.36%.",
        "",
        "## Slide 7: Bias/R2 Tradeoff",
        "- Gated correction has near-zero weighted bias and highest weighted R2.",
        "",
        "## Slide 8: Household Accuracy",
        "- Gated MAE is about 2.47 trips per household.",
        "- 54.5% of households are within 2 trips.",
        "",
        "## Slide 9: Diagnostics",
        "- Global and random controls show that event-level downscaling is the main effect.",
        "- LLM cohort ranking adds smaller but measurable subgroup signal.",
        "",
        "## Slide 10: Takeaway",
        "- LLMs are useful here as event-prior generators, not direct numerical predictors.",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def add_title(slide: Slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.55))
    text_frame = title_box.text_frame
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(15, 23, 42)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.58), Inches(0.92), Inches(12), Inches(0.36))
        sub_frame = sub_box.text_frame
        sub_frame.text = subtitle
        sub_frame.paragraphs[0].font.name = "Microsoft YaHei"
        sub_frame.paragraphs[0].font.size = Pt(13)
        sub_frame.paragraphs[0].font.color.rgb = RGBColor(71, 85, 105)


def add_footer(slide: Slide, index: int) -> None:
    footer = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(12.1), Inches(0.22))
    frame = footer.text_frame
    frame.text = f"NHTS 2022 Event Adaptation | {index}"
    paragraph = frame.paragraphs[0]
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = RGBColor(100, 116, 139)


def add_bullets(
    slide: Slide,
    bullets: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    size: int = 18,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 0
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor(30, 41, 59)
        paragraph.space_after = Pt(8)


def add_metric_card(slide: Slide, x: float, y: float, title: str, value: str, note: str, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.7), Inches(1.25))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
    shape.line.color.rgb = color
    shape.line.width = Pt(1.5)
    frame = shape.text_frame
    frame.clear()
    p1 = frame.paragraphs[0]
    p1.text = title
    p1.font.name = "Microsoft YaHei"
    p1.font.size = Pt(10)
    p1.font.color.rgb = RGBColor(71, 85, 105)
    p2 = frame.add_paragraph()
    p2.text = value
    p2.font.name = "Microsoft YaHei"
    p2.font.size = Pt(23)
    p2.font.bold = True
    p2.font.color.rgb = color
    p3 = frame.add_paragraph()
    p3.text = note
    p3.font.name = "Microsoft YaHei"
    p3.font.size = Pt(8)
    p3.font.color.rgb = RGBColor(100, 116, 139)


def add_picture(slide: Slide, image_path: Path, left: float, top: float, width: float) -> None:
    slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), width=Inches(width))


def create_presentation(summary: pd.DataFrame, accuracy: pd.DataFrame, figure_paths: dict[str, Path]) -> Path:
    path = OUTPUT_DIR / "NHTS_LLM_Event_Adaptation_Presentation.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    blue = RGBColor(59, 130, 246)
    green = RGBColor(22, 163, 74)
    orange = RGBColor(249, 115, 22)

    def slide_base(title: str, subtitle: str | None = None) -> Slide:
        slide = prs.slides.add_slide(blank)
        background = slide.background
        background.fill.solid()
        background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        add_title(slide, title, subtitle)
        add_footer(slide, len(prs.slides))
        return slide

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(15, 23, 42)
    title = slide.shapes.add_textbox(Inches(0.75), Inches(1.15), Inches(11.9), Inches(1.3))
    tf = title.text_frame
    tf.text = "Label-Free LLM Event Adaptation"
    tf.paragraphs[0].font.name = "Microsoft YaHei"
    tf.paragraphs[0].font.size = Pt(42)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    subtitle = slide.shapes.add_textbox(Inches(0.78), Inches(2.55), Inches(10.9), Inches(0.8))
    subtitle.text_frame.text = "Post-pandemic household travel prediction with no 2022 training labels"
    subtitle.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    subtitle.text_frame.paragraphs[0].font.size = Pt(22)
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(203, 213, 225)
    add_metric_card(slide, 0.85, 4.55, "Weighted MAE", "-42.3%", "primary gated rule", blue)
    add_metric_card(slide, 3.85, 4.55, "Weighted Bias", "-99.4%", "gated absolute bias reduction", green)
    add_metric_card(slide, 6.85, 4.55, "LLM Requests", "-83.2%", "cohort-level generation", orange)
    add_footer(slide, 1)

    slide = slide_base("Problem: 2022 breaks routine mobility continuity")
    add_bullets(
        slide,
        [
            "Pre-pandemic historical waves learn routine household mobility.",
            "COVID-era recovery changed commuting, shopping, transit use, and discretionary activity.",
            "The target is household daily trip count, evaluated on 7,893 2022 households.",
        ],
        0.8,
        1.45,
        5.4,
        4.5,
        20,
    )
    add_picture(slide, figure_paths["metric_comparison"], 6.35, 1.35, 6.35)

    slide = slide_base("Research question", "Can event semantics repair 2022 prediction without target-year labels?")
    add_bullets(
        slide,
        [
            "No 2022 CNTTDHH labels are used for training or calibration.",
            "LLM is not used as a direct numerical predictor.",
            "LLM generates structured pandemic-response priors for household cohorts.",
        ],
        0.95,
        1.55,
        5.7,
        4.2,
        22,
    )
    add_picture(slide, figure_paths["workflow"], 6.25, 1.55, 6.3)

    slide = slide_base("Method overview", "Distill LLM event semantics into a lightweight correction rule")
    add_picture(slide, figure_paths["workflow"], 0.85, 1.35, 11.7)
    add_bullets(
        slide,
        ["Final rule: prediction = routine prediction × clip(1 - alpha × event pressure, min_factor, 1)"],
        1.0,
        6.0,
        11.0,
        0.5,
        17,
    )

    slide = slide_base("LLM event priors", "Cohort-level generation improves cost and auditability")
    add_bullets(
        slide,
        [
            "1,327 cohort prompts cover 7,893 households.",
            "Request reduction: 83.2%, about 5.95x fewer calls.",
            "Primary useful feature: trip_suppression_risk.",
        ],
        0.8,
        1.35,
        5.2,
        4.8,
        21,
    )
    add_picture(slide, figure_paths["pressure_distribution"], 6.25, 1.3, 6.35)

    slide = slide_base("Main result", "Fixed gated LLM event prior sharply reduces household prediction error")
    add_picture(slide, figure_paths["metric_comparison"], 0.8, 1.25, 11.9)

    slide = slide_base("Bias and distributional fit", "Gated correction is nearly unbiased")
    add_picture(slide, figure_paths["bias_r2"], 0.9, 1.25, 11.5)

    slide = slide_base("Household-level accuracy", "Interpret regression results with tolerance bands")
    add_picture(slide, figure_paths["household_accuracy"], 0.85, 1.25, 11.7)
    gated_acc = accuracy[accuracy["method"] == "gated_trip_suppression_a1_d0p15"].iloc[0]
    add_bullets(
        slide,
        [
            f"Gated correction: MAE {gated_acc.unweighted_mae:.2f}; "
            f"{pct(gated_acc.within_2_trips)} within ±2 trips; "
            f"{pct(gated_acc.within_3_trips)} within ±3 trips."
        ],
        1.2,
        6.25,
        10.5,
        0.4,
        16,
    )

    slide = slide_base("Diagnostics", "Event-level downscaling is the main effect; cohort ranking adds local signal")
    add_picture(slide, figure_paths["subgroup_gains"], 0.85, 1.25, 7.0)
    add_bullets(
        slide,
        [
            "Global controls are strong: the pandemic shift is a broad event-level effect.",
            "Random controls are weaker than structured event adaptation.",
            "Subgroup gains show where cohort-specific ranking adds value.",
        ],
        8.1,
        1.55,
        4.2,
        4.8,
        19,
    )

    slide = slide_base("Takeaway", "LLMs are useful as event-prior generators, not direct trip-count predictors")
    add_bullets(
        slide,
        [
            "Contribution: label-free event adaptation for post-pandemic travel demand.",
            "Accuracy: primary weighted MAE reduction 42.3%; near-zero weighted bias.",
            "Bias: primary gated correction reduces absolute weighted bias by 99.4%.",
            "Efficiency: cohort prompting reduces LLM requests by 83.2%.",
            "Future work: stronger priors, external event context, and prospective validation.",
        ],
        1.1,
        1.55,
        10.8,
        4.8,
        23,
    )

    prs.save(path)
    return path


def main() -> None:
    configure_logging()
    ensure_dirs()
    set_plot_style()
    metrics = load_label_free_metrics()
    summary = aggregate_metrics(metrics)
    frame_2022, predictions = load_2022_predictions()
    accuracy = build_household_accuracy(frame_2022, predictions)

    metrics_path = OUTPUT_DIR / "final_metrics_summary.csv"
    accuracy_path = OUTPUT_DIR / "household_accuracy_summary.csv"
    save_csv(summary, metrics_path)
    save_csv(accuracy, accuracy_path)

    figure_paths = {
        "metric_comparison": save_metric_comparison(summary),
        "bias_r2": save_bias_r2_tradeoff(summary),
        "household_accuracy": save_household_accuracy(accuracy),
        "pressure_distribution": save_pressure_distribution(frame_2022),
        "subgroup_gains": save_subgroup_gains(),
        "workflow": save_workflow_diagram(),
    }
    report_path = write_report(summary, accuracy, figure_paths)
    storyboard_path = write_storyboard()
    pptx_path = create_presentation(summary, accuracy, figure_paths)

    assets = GeneratedAssets(
        report_path=report_path,
        storyboard_path=storyboard_path,
        pptx_path=pptx_path,
        metrics_path=metrics_path,
        accuracy_path=accuracy_path,
    )
    LOGGER.info("Generated final assets: %s", assets)


if __name__ == "__main__":
    main()
