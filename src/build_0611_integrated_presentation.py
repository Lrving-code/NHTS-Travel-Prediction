"""Build the integrated 0611-based presentation deck."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

import build_template_presentation as base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REFS_DIR = PROJECT_ROOT / "refs"
OUTPUT_PATH = PROJECT_ROOT / "outputs" / "final_project" / "NHTS_Travel_Behavior_0611_Integrated_Presentation.pptx"
NOTES_PATH = PROJECT_ROOT / "outputs" / "final_project" / "NHTS_Travel_Behavior_0611_Integrated_Speaker_Notes.md"
MAIN_SLIDES = 25

LOGGER = logging.getLogger(__name__)


def configure_logging() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")


def source_pptx() -> Path:
    matches = sorted(REFS_DIR.glob("0611*.pptx"), key=lambda path: path.stat().st_size, reverse=True)
    for path in matches:
        if "汇报" in path.name:
            return path
    if matches:
        return matches[0]
    raise FileNotFoundError("No refs/0611*.pptx source deck was found.")


def extract_logo_from(prs: Presentation) -> bytes | None:
    for slide in prs.slides:
        for shape in base.iter_nested_shapes(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            width = shape.width / 914400
            height = shape.height / 914400
            size = len(shape.image.blob)
            if 0.45 <= width <= 1.6 and 0.25 <= height <= 1.0 and size < 120_000:
                return shape.image.blob
    return None


def create_prs_from_source(source: Path) -> tuple[Presentation, bytes | None]:
    prs = Presentation(source)
    logo = extract_logo_from(prs)
    base.clear_template_slides(prs)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    base.TOTAL_SLIDES = MAIN_SLIDES
    return prs, logo


def build_values() -> tuple[dict[str, float], pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    trip, acc, mode, mode_trips, purpose, ci, paired_ci = base.load_results()
    strong = base.read_csv_or_empty(PROJECT_ROOT / "outputs" / "strong_baselines" / "strong_tabular_baseline_metrics.csv")
    baseline = base.metric(trip, "historical_xgboost", "weighted_mae")
    gated = base.metric(trip, "gated_trip_suppression_a1_d0p15", "weighted_mae")
    catboost = float("nan")
    catboost_bias = float("nan")
    catboost_r2 = float("nan")
    if not strong.empty and "catboost_gpu" in set(strong["method"]):
        row = strong[strong["method"] == "catboost_gpu"].iloc[0]
        catboost = float(row["weighted_mae"])
        catboost_bias = float(row["weighted_bias"])
        catboost_r2 = float(row["weighted_r2"])
    mode_trip_base = float("nan")
    mode_trip_hybrid = float("nan")
    if not mode_trips.empty:
        mode_trip_base = float(mode_trips.loc[mode_trips["method"] == "traditional_count_x_traditional_mode", "weighted_total_mode_trip_mae"].iloc[0])
        mode_trip_hybrid = float(mode_trips.loc[mode_trips["method"] == "gated_count_x_llm_mode", "weighted_total_mode_trip_mae"].iloc[0])
    main = "gated_trip_suppression_a1_d0p15"
    values = {
        "baseline": baseline,
        "baseline_bias": base.metric(trip, "historical_xgboost", "weighted_bias"),
        "baseline_r2": base.metric(trip, "historical_xgboost", "weighted_r2"),
        "catboost": catboost,
        "catboost_bias": catboost_bias,
        "catboost_r2": catboost_r2,
        "gated": gated,
        "gated_bias": base.metric(trip, main, "weighted_bias"),
        "gated_r2": base.metric(trip, main, "weighted_r2"),
        "llm_only": base.metric(trip, "llm_only_trip_suppression_a1p25", "weighted_mae"),
        "llm_only_bias": base.metric(trip, "llm_only_trip_suppression_a1p25", "weighted_bias"),
        "llm_only_r2": base.metric(trip, "llm_only_trip_suppression_a1p25", "weighted_r2"),
        "zero_rule": base.metric(trip, "zero_shot_llm_rule_tree", "weighted_mae"),
        "zero_rule_bias": base.metric(trip, "zero_shot_llm_rule_tree", "weighted_bias"),
        "zero_rule_r2": base.metric(trip, "zero_shot_llm_rule_tree", "weighted_r2"),
        "small_rule": base.metric(trip, "llm_rule_small_hist_calibrated_n500", "weighted_mae"),
        "small_rule_bias": base.metric(trip, "llm_rule_small_hist_calibrated_n500", "weighted_bias"),
        "small_rule_r2": base.metric(trip, "llm_rule_small_hist_calibrated_n500", "weighted_r2"),
        "mean_bias": base.metric(trip, "historical_mean_only", "weighted_bias"),
        "trend_bias": base.metric(trip, "historical_mean_trend_shift", "weighted_bias"),
        "mae_reduction": base.pct_delta(baseline, gated),
        "exact": base.metric(acc, main, "exact_rounded_accuracy"),
        "within2": base.metric(acc, main, "within_2_trips"),
        "within3": base.metric(acc, main, "within_3_trips"),
        "transit_base": base.metric(mode, "historical_xgboost", "transit_share_weighted_mae"),
        "transit_llm": base.metric(mode, "llm_transit_avoidance_a1", "transit_share_weighted_mae"),
        "transit_gain": base.pct_delta(base.metric(mode, "historical_xgboost", "transit_share_weighted_mae"), base.metric(mode, "llm_transit_avoidance_a1", "transit_share_weighted_mae")),
        "mode_tv_base": base.metric(mode, "historical_xgboost", "weighted_total_variation"),
        "mode_tv_llm": base.metric(mode, "llm_transit_avoidance_a1", "weighted_total_variation"),
        "mode_trip_base": mode_trip_base,
        "mode_trip_hybrid": mode_trip_hybrid,
        "mode_trip_gain": base.pct_delta(mode_trip_base, mode_trip_hybrid) if mode_trip_base == mode_trip_base else float("nan"),
        "cohort_global_delta": base.cohort_value("overall_primary_vs_global_a1_mae_delta"),
        "cohort_global_a125_delta": base.cohort_value("overall_primary_vs_global_a1p25_mae_delta"),
        "cohort_subgroup_share": base.cohort_value("share_subgroup_cells_primary_beats_global_a1"),
        "cohort_gate_weighted": base.cohort_value("gated_uses_cohort_prior_share_weighted"),
    }
    return values, trip, mode, mode_trips, purpose, ci, paired_ci


def add_0611_title(prs: Presentation, logo: bytes | None, values: dict[str, float]) -> None:
    slide = base.blank_slide(prs)
    base.set_background(slide, base.COLORS["navy"])
    base.add_logo(slide, logo, 11.85, 0.34, 0.92, True)
    base.text_box(slide, "面向时序迁移的家庭出行行为预测", 0.74, 0.8, 10.2, 0.55, 25, base.COLORS["white"], True)
    base.text_box(slide, "大模型事件先验与历史模型校准", 0.74, 1.45, 10.1, 0.48, 21, base.RGBColor(226, 232, 240), True)
    base.text_box(slide, "LLM event priors + historically calibrated prediction under target-year shift", 0.78, 2.12, 9.9, 0.3, 11.5, base.RGBColor(203, 213, 225), True)
    for idx, (title, note) in enumerate([
        ("冷启动选择器", "LLM-rule selector"),
        ("时序迁移修正", "LLM-corrected XGBoost"),
        ("因果/泄露护栏", "causal guardrails"),
        ("多目标评估", "accuracy / bias / equity"),
    ]):
        base.node(slide, 0.78 + idx * 3.0, 3.25, 2.28, 0.72, title, note, base.COLORS["teal"], base.RGBColor(30, 64, 91))
        if idx < 3:
            base.arrow(slide, 3.1 + idx * 3.0, 3.5, 0.48, 0.18, base.RGBColor(148, 163, 184))
    base.metric_card(slide, 0.88, 5.05, 2.55, 0.88, "Trip-count wMAE", f"{values['gated']:.3f}", f"-{values['mae_reduction']:.1%} vs XGBoost", base.COLORS["blue"], base.COLORS["white"])
    base.metric_card(slide, 3.75, 5.05, 2.55, 0.88, "Weighted bias", f"{values['gated_bias']:+.3f}", "nearly unbiased", base.COLORS["green"], base.COLORS["white"])
    base.metric_card(slide, 6.62, 5.05, 2.55, 0.88, "LLM calls", "89", "batch prior prompts", base.COLORS["orange"], base.COLORS["white"])
    base.metric_card(slide, 9.49, 5.05, 2.55, 0.88, "Outputs", "3+", "trips / mode / purpose", base.COLORS["purple"], base.COLORS["white"])


def add_0611_agenda(prs: Presentation, logo: bytes | None) -> None:
    slide = base.blank_slide(prs)
    base.set_background(slide)
    base.add_frame(slide, "00", "目录 / Contents", "按一条主线串起 selector branch 与 correction branch", 2, logo)
    rows = [
        ["01", "研究背景与相关工作", "为什么目标年迁移需要 LLM event-generalization"],
        ["02", "问题定义与数据", "NHTS household behavior system：出行多少、方式、目的"],
        ["03", "方案设计", "冷启动选择器 + LLM 修正 XGBoost 的统一框架"],
        ["04", "性能评估", "传统基线、纯 LLM、rule tree、hybrid adapter 多指标比较"],
        ["05", "结论与展望", "因果护栏、外部验证、本地开源 LLM 控制"],
    ]
    base.small_table(slide, ["章节", "内容", "本页以后怎么讲"], rows, 0.82, 1.35, [1.0, 3.0, 7.0], 0.72, 12, base.COLORS["navy2"])
    base.story_box(slide, "主线不是“两个方案拼接”", "The selector branch handles cold-start simulation; the correction branch handles target-year temporal adaptation.", 1.0, 5.78, 10.7, 0.78, base.COLORS["teal"])


def add_two_scenarios_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = base.blank_slide(prs)
    base.set_background(slide)
    base.add_frame(slide, "03", "问题定义 / Problem Definition", "两个场景对应同一个 event-generalization 框架", 6, logo)
    base.story_box(slide, "场景 A：数据稀疏 / 字段缺失 / 冷启动", "Historical labels are weak or missing. LLM knowledge is distilled into fast executable rules; low-confidence cases can fall back to LLM.", 0.72, 1.25, 5.65, 1.35, base.COLORS["orange"])
    base.story_box(slide, "场景 B：有历史数据但目标年情境变化", "Historical NHTS labels exist, but 2022 context changes the meaning of routine covariates. LLM event priors correct the historical predictor.", 6.78, 1.25, 5.65, 1.35, base.COLORS["teal"])
    base.node(slide, 1.0, 3.25, 2.1, 0.65, "Pure data", "calibrated but brittle", base.COLORS["blue"], base.COLORS["light"])
    base.arrow(slide, 3.2, 3.48, 0.55, 0.16, base.COLORS["gray"])
    base.node(slide, 3.85, 3.25, 2.1, 0.65, "Pure LLM", "general but unstable", base.COLORS["purple"], base.COLORS["light"])
    base.arrow(slide, 6.05, 3.48, 0.55, 0.16, base.COLORS["gray"])
    base.node(slide, 6.72, 3.25, 2.15, 0.65, "Selector", "cold-start branch", base.COLORS["orange"], base.COLORS["light"])
    base.arrow(slide, 8.96, 3.48, 0.55, 0.16, base.COLORS["gray"])
    base.node(slide, 9.63, 3.25, 2.15, 0.65, "Correction", "temporal-adaptation branch", base.COLORS["green"], base.COLORS["light"])
    base.small_table(
        slide,
        ["模块", "输入", "LLM 做什么", "输出"],
        [
            ["Selector branch", "少量/缺失历史样本", "蒸馏 IF-THEN 规则与置信度", "快速模拟预测 + 低置信回退"],
            ["Correction branch", "2001/2009/2017 历史 + 2022 无标签 profiles", "生成 trip suppression / remote work / transit avoidance priors", "修正 XGBoost 的 2022 household prediction"],
        ],
        0.92,
        4.52,
        [2.0, 3.0, 4.05, 2.8],
        0.56,
        12,
        base.COLORS["navy2"],
    )


def add_selector_branch_slide(prs: Presentation, logo: bytes | None) -> None:
    slide = base.blank_slide(prs)
    base.set_background(slide)
    base.add_frame(slide, "06", "冷启动选择器 / Selector Branch", "把同学方案纳入统一框架：LLM 不是逐条替代模型，而是蒸馏预测路径", 8, logo)
    stages = [
        ("LLM rule extraction", "从大模型提取可执行规则"),
        ("Rule prediction", "常规样本快速预测"),
        ("Confidence scoring", "识别低置信或冲突样本"),
        ("LLM fallback", "必要时再调用大模型"),
    ]
    for idx, (title, note) in enumerate(stages):
        x = 0.78 + idx * 3.02
        color = [base.COLORS["purple"], base.COLORS["blue"], base.COLORS["orange"], base.COLORS["teal"]][idx]
        base.node(slide, x, 1.5, 2.3, 0.78, title, note, color, base.COLORS["light"])
        if idx < 3:
            base.arrow(slide, x + 2.36, 1.78, 0.48, 0.18, base.COLORS["gray"])
    base.small_table(
        slide,
        ["实验场景", "为什么需要 selector", "证明什么"],
        [
            ["数据稀疏", "历史样本只有 100/500/1000 条", "LLM priors 可缓解 cold-start"],
            ["字段缺失", "随机缺失 30%-70% 字段", "规则/低置信机制比纯数据拟合更稳"],
            ["新增因素", "出现历史数据没有的新情境", "需要 event-generalization 而不是只靠历史拟合"],
        ],
        0.85,
        3.1,
        [2.2, 4.15, 4.45],
        0.58,
        12,
        base.COLORS["navy2"],
    )
    base.story_box(slide, "它和我们的主方法如何统一？", "Selector branch solves when to trust distilled rules or call an LLM; correction branch solves how to numerically calibrate household travel under a known target-year shift.", 1.0, 5.75, 10.8, 0.82, base.COLORS["green"])


def add_local_llm_backup(prs: Presentation, logo: bytes | None) -> None:
    base.add_backup_qa_slide(
        prs,
        logo,
        "B10",
        "开源/本地 LLM 复刻做到哪一步？",
        "What is the status of local open-source LLM replication?",
        "脚本和环境审计已经实现；当前 Python 环境是 CPU-only PyTorch，不能声称已完成 GPU local run。",
        [
            ["GPU hardware", "RTX 4090 visible through nvidia-smi"],
            ["Current blocker", "PyTorch CUDA available = False"],
            ["Implemented artifact", "src/run_local_llm_prior_replication.py"],
            ["Paper-safe wording", "protocol + audit implemented; generation pending"],
        ],
        "答法：这是我们下一步最值得补的实验。只要换成 CUDA-enabled PyTorch，并准备本地 Qwen/Llama 类 instruct model，就能在 32/64 个 cohort 上生成 local priors，再和 GPT-5.5 prior 做 Spearman/Pearson 排序一致性比较。",
    )


def apply_text_replacements(prs: Presentation) -> None:
    replacements = {
        "蒸馏预测路径": "提取预测路径",
        "LLM 蒸馏与部署": "LLM 先验生成与部署",
        "Distillation & Deployment": "Prior Generation & Deployment",
        "Rule distillation": "Rule extraction",
        "rule distillation": "rule extraction",
        "规则蒸馏和低置信": "规则提取和低置信",
        "How do rule distillation": "How do rule extraction",
        "蒸馏 IF-THEN": "提取 IF-THEN",
        "蒸馏出可执行": "提取可执行",
        "从逐户调用转成 batch prior、规则适配和置信度回退": "从逐户调用转成 batch prior、规则提取和置信度回退",
    }
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            text = shape.text_frame.text
            updated = text
            for old, new in replacements.items():
                updated = updated.replace(old, new)
            if updated != text:
                shape.text_frame.text = updated


def normalize_main_slide_numbers(prs: Presentation) -> None:
    section_labels = ["TITLE", "00", *[f"{index:02d}" for index in range(1, MAIN_SLIDES - 2)], "END"]
    for slide_index, slide in enumerate(list(prs.slides)[:MAIN_SLIDES], start=1):
        label = section_labels[slide_index - 1]
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            left = shape.left / 914400
            top = shape.top / 914400
            text = shape.text_frame.text.strip()
            if 0.25 <= top <= 0.7 and 0.2 <= left <= 1.1 and text:
                shape.text_frame.text = label
            if 6.85 <= top <= 7.25 and left >= 10.7 and "/" in text:
                shape.text_frame.text = f"{slide_index:02d} / {MAIN_SLIDES}"


def add_integrated_slides(prs: Presentation, logo: bytes | None) -> None:
    values, trip, mode, _mode_trips, purpose, ci, paired_ci = build_values()
    add_0611_title(prs, logo, values)
    add_0611_agenda(prs, logo)
    base.add_problem_slide(prs, logo, values)
    base.add_related_work_slide(prs, logo)
    base.add_data_slide(prs, logo)
    add_two_scenarios_slide(prs, logo)
    base.add_method_spectrum_slide(prs, logo, values)
    add_selector_branch_slide(prs, logo)
    base.add_route_slide(prs, logo)
    base.add_llm_prior_slide(prs, logo)
    base.add_comparison_slide(prs, logo, trip)
    base.add_results_slide(prs, logo, trip, values)
    base.add_statistical_validation_slide(prs, logo, ci, paired_ci)
    base.add_tradeoff_slide(prs, logo)
    base.add_multi_objective_slide(prs, logo)
    base.add_error_distribution_figure_slide(prs, logo)
    base.add_event_heterogeneity_slide(prs, logo)
    base.add_robustness_slide(prs, logo)
    base.add_subgroup_slide(prs, logo)
    base.add_error_insight_slide(prs, logo, values)
    base.add_llm_role_slide(prs, logo, values)
    base.add_mode_slide(prs, logo, mode, values)
    base.add_purpose_slide(prs, logo, purpose)
    base.add_scalability_slide(prs, logo)
    base.add_final_slide(prs, logo, values)
    base.add_backup_qa_slides(prs, logo, values)
    base.add_backup_cohort_value_slide(prs, logo, values)
    add_local_llm_backup(prs, logo)
    apply_text_replacements(prs)
    normalize_main_slide_numbers(prs)



def write_notes() -> None:
    lines = [
        "# 0611 Integrated Speaker Notes",
        "",
        "## 10 分钟主讲路径",
        "",
        "1. 第 1-5 页：说明问题不是普通出行预测，而是目标年时序迁移；用 2025-2026 文献调研引出 LLM event-generalization。",
        "2. 第 6-8 页：把两方工作统一成两条分支：selector branch 处理冷启动/字段缺失，correction branch 处理有历史数据但目标年情境变化。",
        "3. 第 9-12 页：讲主方法。历史 XGBoost 学 routine mobility，LLM 生成 trip suppression / remote work / transit avoidance 等结构化 priors，再做 no-label correction。",
        "4. 第 13-19 页：讲证据链。主结果、bootstrap CI、误差-偏差权衡、多目标选择、随机置换、异质性和分组稳健性。",
        "5. 第 20-25 页：解释 LLM 的真实角色、mode/purpose 扩展、batch prompting/部署逻辑和总结。",
        "",
        "## 关键口径",
        "",
        "- 不要说 LLM 直接预测家庭出行次数；要说 LLM 生成 event prior / selector signal。",
        "- 不要说两个方案拼接；要说它们是同一个 temporal-adaptation framework 的两个 operating regimes。",
        "- 不要把总标题写成“知识蒸馏”或“模型蒸馏”；本项目更准确的说法是规则提取、事件先验生成和历史模型校准。",
        "- 不要隐瞒 global prior 很强；应说主贡献是 event-level correction，cohort ranking 是 selective refinement。",
        "- local/open-source LLM replication 目前是 protocol + environment audit，当前环境被 CPU-only PyTorch 阻塞，不能作为已完成结果。",
    ]
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def build_deck() -> Path:
    source = source_pptx()
    prs, logo = create_prs_from_source(source)
    add_integrated_slides(prs, logo)
    path = base.save_presentation(prs, OUTPUT_PATH)
    write_notes()
    return path


def main() -> None:
    configure_logging()
    path = build_deck()
    LOGGER.info("Wrote integrated 0611-based presentation: %s", path)


if __name__ == "__main__":
    main()
